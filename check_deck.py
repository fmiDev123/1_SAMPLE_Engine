"""Is this exported deck missing any of its numbers?

`check_workbook.py` asks whether a WORKBOOK can carry a report.  This asks the
question at the other end: having built the deck, did every exhibit actually
get its data?

It reports blank exhibits - a chart whose series are empty or all zero, a table
with no body rows or a body that is all zeros and dashes - and groups them by
slide TITLE, so one template failing across every country shows up as a single
line with a count rather than six hundred near-identical ones.  That grouping
is the point: it turns "some slides are blank" into "16 tables x 38 countries,
all from one lookup", which is a cause rather than a symptom.

    python check_deck.py "1 - Some Market Report (FMI).pptx"

Add a number to widen the example-slide list (default 3):

    python check_deck.py DECK.pptx 10

Exits 1 when anything is flagged, so it can gate a release.

WHAT IT CANNOT SEE
------------------
A page the registry DROPPED leaves no trace in the file, so a whole missing
section reads as a clean deck.  That is how a workbook whose Unit column said
'Volume (MT)' lost its entire volume and pricing sections silently.  To catch
that, compare slide counts against a deck you trust:

    python check_deck.py GOOD.pptx | head -2
    python check_deck.py SUSPECT.pptx | head -2

A deck that is much SHORTER than a comparable one is the tell.

KNOWN FALSE POSITIVES
---------------------
Bubble charts legitimately carry one point per series, so they are not
flagged.  Judge anything else against a deck known to be good before treating
it as a defect.
"""
from __future__ import annotations

import re
import sys
import threading
from collections import defaultdict

# python-pptx walks package relationships recursively, and a 1,000-slide deck
# overflows the default stack while merely OPENING the file.  The work is done
# on a thread with a bigger stack (256 MB is refused on Windows; 64 works).
sys.setrecursionlimit(60_000)
threading.stack_size(64 * 1024 * 1024)

GROUP = 6                       # MSO_SHAPE_TYPE.GROUP
DEAD_TEXT = {"", "-", "–", "—", "--", "n/a", "na"}
NUMISH = re.compile(r"^[\s\-–—]*[\d,.]*[\s%$]*$")


def _shapes(shapes):
    """Every shape, including the ones nested inside groups - most of this
    deck's charts and tables live in groups, so a flat loop misses them."""
    for sh in shapes:
        yield sh
        if sh.shape_type == GROUP:
            yield from _shapes(sh.shapes)


def _dead_cell(text: str) -> bool:
    """True for a body cell carrying no information: blank, a dash, or a
    number that is entirely zero ('0', '0.0%', '0.00')."""
    s = text.strip()
    if s.casefold() in DEAD_TEXT:
        return True
    if NUMISH.match(s):
        digits = re.sub(r"[^\d]", "", s)
        return bool(digits) and int(digits) == 0
    return False


def audit(path: str, show: int = 3) -> int:
    from pptx import Presentation

    prs = Presentation(path)
    found = defaultdict(list)

    for idx, slide in enumerate(prs.slides, start=1):
        texts = [sh.text_frame.text.strip() for sh in _shapes(slide.shapes)
                 if sh.has_text_frame and sh.text_frame.text.strip()]
        title = re.sub(r"\s+", " ", texts[0])[:72] if texts else "(untitled)"

        for sh in _shapes(slide.shapes):
            if getattr(sh, "has_chart", False):
                chart = sh.chart
                values = [v for s in chart.series for v in s.values
                          if v is not None]
                if not values:
                    found[("chart-empty", title)].append(idx)
                elif not any(values):
                    found[("chart-all-zero", title)].append(idx)

            if getattr(sh, "has_table", False):
                body = list(sh.table.rows)[1:]          # row 0 is the header
                if not body:
                    found[("table-no-rows", title)].append(idx)
                    continue
                # column 0 holds row LABELS, which are text even when the
                # figures beside them are missing - so judge only the body.
                cells = [c.text for r in body for c in list(r.cells)[1:]]
                if cells and all(_dead_cell(c) for c in cells):
                    found[("table-all-blank", title)].append(idx)

    print(f"deck   : {path}")
    print(f"slides : {len(prs.slides)}")
    print()
    if not found:
        print("PASS  every chart and table carries data.")
        return 0

    total = sum(len(v) for v in found.values())
    print(f"{'kind':18s} {'n':>5s}  slide title")
    print("-" * 96)
    for (kind, title), slides in sorted(found.items(), key=lambda kv: -len(kv[1])):
        print(f"{kind:18s} {len(slides):5d}  {title}")
        more = " ..." if len(slides) > show else ""
        print(f"{'':25s}slides: {slides[:show]}{more}")
    print()
    print(f"FAIL  {total} blank exhibits across {len(found)} distinct pages.")
    print("      A count equal to (pages x geographies) points at ONE lookup,")
    print("      not at the data - check the Unit and Segment column spellings")
    print("      in the workbook before suspecting the figures.")
    return 1


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    show = int(sys.argv[2]) if len(sys.argv) > 2 else 3
    return audit(sys.argv[1], show)


if __name__ == "__main__":
    code = []
    t = threading.Thread(target=lambda: code.append(main()))
    t.start()
    t.join()
    sys.exit(code[0] if code else 1)
