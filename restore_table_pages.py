"""Fill the body-less table pages of a finished deck with their real content.

    python restore_table_pages.py "deck.pptx" "source.pptx" ["out.pptx"]

``deck`` is the deck with title-only table pages ("... Market Tables X (3/6)"
with no table under the title).  ``source`` is a deck built from the same
workbook and settings WITHOUT the volume-removal step, so every page exists
(build it with FMI_DECK_STYLE=off).  For each empty page the missing shapes -
the panel card, its title, the table, the source line, any padlock - are
copied in from the source slide of the same title.  Nothing else in the deck
is touched, so a locked deck stays locked: the copied table is the locked one.
"""
import copy
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation                      # noqa: E402
from fmi_report import deck_style as D             # noqa: E402

R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _sig(sh):
    return (sh.name, sh.left, sh.top, sh.width, sh.height)


def _copy_shape(src_sl, dst_sl, sh):
    el = copy.deepcopy(sh._element)
    for node in el.iter():
        rid = node.get(R + "embed")
        if rid:                                        # a picture: re-attach
            blob = src_sl.part.related_part(rid).blob
            _, new = dst_sl.part.get_or_add_image_part(io.BytesIO(blob))
            node.set(R + "embed", new)
        for att in (R + "link", R + "id"):             # hyperlinks: leave out
            if node.get(att) and node.tag.endswith("hlinkClick"):
                node.getparent().remove(node)
                break
    dst_sl.shapes._spTree.append(el)


def main(deck, source, out):
    prs, src = Presentation(str(deck)), Presentation(str(source))
    by_title = {}
    for sl in src.slides:
        by_title.setdefault(D.title_of(sl), []).append(sl)
    filled, missing = 0, []
    for i, sl in enumerate(prs.slides, 1):
        t = D.title_of(sl)
        if not (D._PAGED.match(t) and D.empty_body(sl)):
            continue
        cands = by_title.get(t)
        if not cands:
            missing.append((i, t))
            continue
        have = {_sig(s) for s in sl.shapes}
        n = 0
        for sh in cands[0].shapes:
            if _sig(sh) not in have:
                _copy_shape(cands[0], sl, sh)
                n += 1
        if n:
            filled += 1
    prs.save(str(out))
    print(f"{Path(deck).name}: {filled} pages filled; "
          f"{len(missing)} with no match in the source")
    for m in missing:
        print("  no source page:", m[0], m[1][:80].encode("ascii", "replace").decode())


if __name__ == "__main__":
    a = sys.argv
    main(a[1], a[2], a[3] if len(a) > 3 else a[1])
