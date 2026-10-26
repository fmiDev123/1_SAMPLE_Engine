"""Repair a finished deck in place: delete the paged table / YoY / segmental
slides that have a title but no body, and renumber the pages that remain.

    python fix_blank_table_pages.py "deck.pptx" [out.pptx]

Nothing else in the deck is touched.  The same repair now runs inside every
build (fmi_report/deck_style.py), so a new deck does not need it.
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation                     # noqa: E402
from fmi_report import deck_style as D            # noqa: E402

src = Path(sys.argv[1])
dst = Path(sys.argv[2]) if len(sys.argv) > 2 else src
prs = Presentation(str(src))
before = len(prs.slides)
gone = D.drop_empty_table_pages(D.Deck(prs))
prs.save(str(dst))
print(f"{src.name}: {gone} empty pages removed ({before} -> {len(prs.slides)} slides)")
