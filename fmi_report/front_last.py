"""Front page and last page of the FMI deck.

The approved artwork lives in ``assets/front_last_template.pptx`` (slide 1 =
front page, slide 2 = "Thank You" / contact page).  Every FMI-template deck -
the unlocked report AND the locked sample - gets:

  * slide 1 replaced by the front page, filled with this deck's market,
    geography, years and "Sample Report" / "Report" tag;
  * the contact page appended as the last slide.

It works on the finished .pptx (like deck_style), so it is the last thing a
deck goes through and nothing in the lock / style passes can touch it.  The
template's layout art (background photo, blue panel, yellow band) is copied
onto the slide itself and its theme colours are written out as RGB, so the
page looks the same whatever theme the deck is built on.

Fact.MR decks (2026-10-14) use their own approved set from
``assets/factmr_front_last_template.pptx``: slide 1 = front page, slides 2 and 3
= the about / "Thank you for trusting Fact.MR" page and the contact page, both
appended at the end.
Switch off for a build with FMI_FRONT_LAST=off.
"""
from __future__ import annotations

import copy
import io
import os
from datetime import datetime
from pathlib import Path

from lxml import etree
from pptx import Presentation

_ASSETS = Path(__file__).resolve().parent / "assets"
_TEMPLATE = _ASSETS / "front_last_template.pptx"            # FMI
_TEMPLATE_TMR = _ASSETS / "tmr_front_last_template.pptx"    # Transparency MR
_TEMPLATE_FACTMR = _ASSETS / "factmr_front_last_template.pptx"   # Fact.MR
_NS = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
}
_A = "{%s}" % _NS["a"]
_R = "{%s}" % _NS["r"]
_P = "{%s}" % _NS["p"]
_SLOT = {"bg1": "lt1", "tx1": "dk1", "bg2": "lt2", "tx2": "dk2"}

_LAST = {"template": None}          # brand of the most recent export


def note_template(name) -> None:
    """export_pptx_native tells us which template it built, so the locked
    path (which runs later, on the saved file) knows whether to apply."""
    _LAST["template"] = name


def enabled() -> bool:
    return os.environ.get("FMI_FRONT_LAST", "on").strip().lower() != "off"


def brand(name=None):
    """'fmi' / 'tmr' / 'factmr' - the templates that carry approved front +
    last pages, else None."""
    from . import theming
    key = theming.normalise(name or _LAST["template"])
    return key if key in ("fmi", "tmr", "factmr") else None


def is_fmi(name=None) -> bool:
    return brand(name) == "fmi"


# ───────────────────────── template helpers ─────────────────────────
def _theme_colours(prs, master=None) -> dict:
    """slot name -> 'RRGGBB' from the template's theme part (the first
    master's, or ``master``'s when given - a template file can carry several
    masters, each on its own theme)."""
    master = master or prs.slide_masters[0]
    theme_part = master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme")
    root = etree.fromstring(theme_part.blob)
    cs = root.find(".//" + _A + "clrScheme")
    out = {}
    for el in cs:
        slot = etree.QName(el).localname
        c = el[0]
        out[slot] = (c.get("lastClr") or c.get("val") or "000000").upper()
    return out


def _resolve_scheme(root, colours) -> None:
    for el in list(root.iter(_A + "schemeClr")):
        val = el.get("val")
        slot = _SLOT.get(val, val)
        rgb = colours.get(slot)
        if rgb is None:
            continue
        new = etree.Element(_A + "srgbClr", val=rgb)
        for ch in el:
            new.append(ch)
        el.getparent().replace(el, new)


def _copy_shapes(src_part_shapes, src_part, dst_slide):
    """Deep-copy shape elements, re-attaching every image they embed to the
    destination slide.  Returns the copied elements."""
    out = []
    for el in src_part_shapes:
        new = copy.deepcopy(el)
        for node in new.iter():
            for attr in ("embed", "link"):
                rid = node.get(_R + attr)
                if not rid:
                    continue
                rel = src_part.rels.get(rid)
                if rel is None or rel.is_external:
                    node.attrib.pop(_R + attr)
                    continue
                _img, new_rid = dst_slide.part.get_or_add_image_part(
                    io.BytesIO(rel.target_part.blob))
                node.set(_R + attr, new_rid)
        out.append(new)
    return out


def _renumber(sp_tree) -> None:
    n = 1
    for c in sp_tree.iter(_P + "cNvPr"):
        n += 1
        c.set("id", str(n))


def _clear(slide) -> None:
    sp_tree = slide.shapes._spTree
    for el in list(sp_tree):
        if etree.QName(el).localname not in ("nvGrpSpPr", "grpSpPr"):
            sp_tree.remove(el)
    for rid, rel in list(slide.part.rels.items()):      # old pictures
        if rel.reltype.endswith("/image"):
            slide.part.drop_rel(rid)


def _paint(slide, tmpl_slide, colours, keep_ph=False) -> list:
    """Flatten the template slide (layout art underneath) onto ``slide``."""
    lay = tmpl_slide.slide_layout
    layer = [e for e in lay.shapes._spTree
             if etree.QName(e).localname in ("sp", "pic", "grpSp", "cxnSp")
             and not e.findall(".//" + _P + "ph")]
    master_layer = []
    if keep_ph:
        # the slide's own master (and its theme) - the footer band, the logo
        mst = lay.slide_master
        colours = _theme_colours(None, mst)
        master_layer = [e for e in mst.shapes._spTree
                        if etree.QName(e).localname in ("sp", "pic", "grpSp",
                                                        "cxnSp")
                        and not e.findall(".//" + _P + "ph")]
    lay_ph = {}
    if keep_ph:
        for e in lay.shapes._spTree:
            ph_ = e.find(".//" + _P + "ph")
            if ph_ is not None:
                lay_ph[ph_.get("idx")] = e
    body = []
    for sh in tmpl_slide.shapes:
        e = sh._element
        if etree.QName(e).localname not in ("sp", "pic", "grpSp", "cxnSp"):
            continue
        ph = e.find(".//" + _P + "ph")
        if ph is None:
            body.append(e)
        elif keep_ph and ph.get("type") not in ("sldNum", "dt"):
            e = copy.deepcopy(e)       # a plain box, at its inherited place
            ph2 = e.find(".//" + _P + "ph")
            ph2.getparent().remove(ph2)
            # keep the layout placeholder's text look (size, colour, anchor)
            lp = lay_ph.get(ph.get("idx"))
            tx_ = e.find(_P + "txBody")
            if lp is not None and tx_ is not None:
                ltx = lp.find(_P + "txBody")
                if ltx is not None:
                    for tag in ("bodyPr", "lstStyle"):
                        mine, theirs = tx_.find(_A + tag), ltx.find(_A + tag)
                        if mine is not None and theirs is not None                                 and len(mine) == 0 and not mine.attrib:
                            new = copy.deepcopy(theirs)
                            tx_.replace(mine, new)
            spPr = e.find(_P + "spPr")
            if spPr is not None and spPr.find(_A + "xfrm") is None:
                x = etree.Element(_A + "xfrm")
                etree.SubElement(x, _A + "off", x=str(int(sh.left)),
                                 y=str(int(sh.top)))
                etree.SubElement(x, _A + "ext", cx=str(int(sh.width)),
                                 cy=str(int(sh.height)))
                spPr.insert(0, x)
            body.append(e)
        elif ph.get("type") in ("ctrTitle", "title", "subTitle"):
            e = copy.deepcopy(e)       # keep it as a plain text box
            ph2 = e.find(".//" + _P + "ph")
            ph2.getparent().remove(ph2)
            body.append(e)
    els = ((_copy_shapes(master_layer, lay.slide_master.part, slide)
            if master_layer else [])
           + _copy_shapes(layer, lay.part, slide)
           + _copy_shapes(body, tmpl_slide.part, slide))
    tree = slide.shapes._spTree
    for el in els:
        _resolve_scheme(el, colours)
        tree.append(el)
    _renumber(tree)
    return els


# ───────────────────────── the front page ─────────────────────────
def _title_size(n: int) -> int:
    # 60 pt fits ~15 characters a line in the 6.1" panel, three lines high
    return 6000 if n <= 36 else 4800 if n <= 60 else 4000 if n <= 90 else 3200


def _set_run_text(par, text) -> None:
    runs = par.findall(_A + "r")
    runs[0].find(_A + "t").text = text
    for r in runs[1:]:
        par.remove(r)


def _fill_front(els, geo_word, market, years, tag) -> None:
    hist0, hist1, est0, last = years
    year = str(datetime.now().year)
    for root in els:
        for sp in root.iter(_P + "sp"):
            nv = sp.find(_P + "nvSpPr/" + _P + "cNvPr")
            name = nv.get("name") if nv is not None else ""
            tx = sp.find(_P + "txBody")
            if tx is None:
                continue
            if name == "Text Placeholder 4":            # title + subtitle
                paras = tx.findall(_A + "p")
                title = next(p for p in paras if p.findall(_A + "r")
                             and "Market" in "".join(
                                 t.text or "" for t in p.iter(_A + "t")))
                sub = next(p for p in paras if "Analysis" in "".join(
                    t.text or "" for t in p.iter(_A + "t")))
                _set_run_text(title, market)
                title.find(_A + "r").find(_A + "rPr").set(
                    "sz", str(_title_size(len(market))))
                _set_run_text(
                    sub, f"{geo_word} Industry Analysis {hist0} – {hist1}"
                         f" and Opportunity Assessment {est0} – {last}")
                # the subtitle gets a box of its own, bottom-anchored just
                # above the yellow band, instead of being pushed down by
                # a stack of empty paragraphs
                sub_sp = copy.deepcopy(sp)
                sub_sp.find(_P + "nvSpPr/" + _P + "cNvPr").set(
                    "name", "Front Page Subtitle")
                for p in tx.findall(_A + "p"):
                    tx.remove(p)
                tx.append(title)
                stx = sub_sp.find(_P + "txBody")
                for p in stx.findall(_A + "p"):
                    stx.remove(p)
                stx.append(sub)
                off = sub_sp.find(".//" + _A + "off")
                ext = sub_sp.find(".//" + _A + "ext")
                off.set("y", "4150000")
                ext.set("cy", "1650000")
                stx.find(_A + "bodyPr").set("anchor", "b")
                sp.addnext(sub_sp)
                ext2 = sp.find(".//" + _A + "ext")
                ext2.set("cy", "3800000")
            elif name == "Rectangle 3":                 # Sample Report| 2026
                runs = tx.findall(".//" + _A + "r")
                runs[0].find(_A + "t").text = tag
                runs[-1].find(_A + "t").text = year
            elif name == "TextBox 9":                   # Source@ ..., 2026
                for t in tx.iter(_A + "t"):
                    t.text = (t.text or "").replace("2026", year)


_BRAND_SUBS = (("Fact.MR Analysis", "TMR Analysis"),
               ("Fact.MR", "TMR"),
               ("Future Market Insights", "Transparency Market Research"),
               ("FMI", "TMR"))


def _rebrand_text(prs) -> None:
    """Text the page builders write with the FMI name (sources, captions,
    the classic write-up's publisher) -> TMR.  Shape names are untouched."""
    import re
    pats = [(re.compile(r"\b%s\b" % re.escape(a)), b) for a, b in _BRAND_SUBS]
    for slide in prs.slides:
        for t in slide._element.iter(_A + "t"):
            if t.text and ("FMI" in t.text or "Future Market" in t.text
                              or "Fact.MR" in t.text):
                for rx, rep in pats:
                    t.text = rx.sub(rep, t.text)


def _rename_shapes(prs, prefix) -> None:
    """Hidden shape names the builders write as 'FMI Section Bar' ... carry
    the FMI name into a TMR / Fact.MR file (2026-10-14: 'no FMI residue').
    Renamed on the finished deck, after the lock pass has used them."""
    for slide in prs.slides:
        for nv in slide._element.iter(_P + "cNvPr"):
            nm = nv.get("name") or ""
            if nm.startswith("FMI "):
                nv.set("name", prefix + nm[4:])


def _year_in(els) -> None:
    year = str(datetime.now().year)
    for root in els:
        for t in root.iter(_A + "t"):
            if t.text and ("©2026" in t.text or "©2025" in t.text):
                t.text = t.text.replace("2026", year).replace("2025", year)


def _fill_front_tmr(els, geo_word, market, years, sample) -> None:
    """Transparency Market Research front page: title + subtitle text boxes
    (the template's placeholders, made plain), the copyright line, and a
    'Sample Report' tag on a locked sample."""
    hist0, hist1, est0, last = years
    _year_in(els)
    for root in els:
        for sp in root.iter(_P + "sp"):
            nv = sp.find(_P + "nvSpPr/" + _P + "cNvPr")
            name = nv.get("name") if nv is not None else ""
            tx = sp.find(_P + "txBody")
            if tx is None:
                continue
            if name == "Title 1":
                _set_run_text(tx.find(_A + "p"), market)
                rpr = tx.find(".//" + _A + "r/" + _A + "rPr")
                n = len(market)
                rpr.set("sz", str(3600 if n <= 40 else 3000 if n <= 60
                                  else 2600))
            elif name == "Subtitle 2":
                _set_run_text(
                    tx.find(_A + "p"),
                    f"{geo_word} Industry Analysis {hist0}–{hist1} and "
                    f"Opportunity Assessment, {est0}–{last}")
                if sample:
                    tag = copy.deepcopy(sp)
                    tag.find(_P + "nvSpPr/" + _P + "cNvPr").set(
                        "name", "Front Page Tag")
                    ttx = tag.find(_P + "txBody")
                    _set_run_text(ttx.find(_A + "p"),
                                  f"Sample Report | {datetime.now().year}")
                    r = ttx.find(".//" + _A + "r/" + _A + "rPr")
                    r.set("b", "1")
                    r.set("sz", "2000")
                    for f in r.findall(_A + "solidFill"):
                        r.remove(f)
                    r.insert(0, etree.fromstring(
                        '<a:solidFill xmlns:a="%s"><a:srgbClr val="0070C0"/>'
                        '</a:solidFill>' % _NS["a"]))
                    tag.find(".//" + _A + "off").set("y", "2450000")
                    sp.addnext(tag)


def _fill_front_factmr(els, geo_word, market, years, sample) -> None:
    """Fact.MR front page: the kicker line becomes the report tag, the title
    is the market, and a subtitle line (geography + study years) goes under
    the three purple bars."""
    hist0, hist1, est0, last = years
    year = str(datetime.now().year)
    for root in els:
        for sp in list(root.iter(_P + "sp")):
            nv = sp.find(_P + "nvSpPr/" + _P + "cNvPr")
            name = nv.get("name") if nv is not None else ""
            tx = sp.find(_P + "txBody")
            if tx is None:
                continue
            if name == "Text Placeholder 7":             # the title
                _set_run_text(tx.find(_A + "p"), market)
                n = len(market)
                rpr = tx.find(".//" + _A + "r/" + _A + "rPr")
                if rpr is not None:
                    rpr.set("sz", str(2800 if n <= 40 else 2400 if n <= 60
                                      else 2000))
                    rpr.set("b", "1")
                # subtitle: its own box under the bars
                sub = copy.deepcopy(sp)
                sub.find(_P + "nvSpPr/" + _P + "cNvPr").set(
                    "name", "Front Page Subtitle")
                stx = sub.find(_P + "txBody")
                _set_run_text(stx.find(_A + "p"),
                              f"{geo_word} Industry Analysis {hist0}\u2013{hist1}"
                              f" and Opportunity Assessment, {est0}\u2013{last}")
                r = stx.find(".//" + _A + "r/" + _A + "rPr")
                if r is not None:
                    r.set("sz", "1800")
                    r.set("b", "0")
                off = sub.find(".//" + _A + "off")
                ext = sub.find(".//" + _A + "ext")
                if off is not None and ext is not None:
                    off.set("y", "3150000")
                    ext.set("cy", "700000")
                sp.addnext(sub)
            elif name == "Footer Placeholder 1" and sp.getparent() is not None \
                    and "RESEARCH" in "".join(t.text or "" for t in tx.iter(_A + "t")):
                _set_run_text(tx.find(_A + "p"),
                              f"{'SAMPLE REPORT' if sample else 'REPORT'}"
                              f"  |  {year}")


# ───────────────────────── entry points ─────────────────────────
def _walk(shapes):
    for sh in shapes:
        try:
            is_group = sh.shape_type == 6
        except NotImplementedError:      # a shape python-pptx cannot classify
            is_group = False
        if is_group:
            yield from _walk(sh.shapes)
        else:
            yield sh


def _old_title(slide, market) -> str:
    """The engine's cover title - the text that ends with the market name."""
    for sh in _walk(slide.shapes):
        if sh.has_text_frame:
            t = sh.text_frame.text.strip()
            if t.endswith(market):
                return t
    return ""


def apply_file(path, sample: bool, template=None) -> bool:
    """Replace slide 1 with the front page and append the last page.
    Returns True when the deck was changed."""
    which = brand(template) if enabled() else None
    tpl_path = {"tmr": _TEMPLATE_TMR, "factmr": _TEMPLATE_FACTMR}.get(
        which, _TEMPLATE)
    if which is None or not tpl_path.exists():
        return False
    from . import config, data
    path = Path(path)
    prs = Presentation(str(path))
    if not len(prs.slides):
        return False
    tpl = Presentation(str(tpl_path))
    colours = _theme_colours(tpl)

    # geography word from the cover the engine drew: "<Geo> <Market>"
    market = str(data.MARKET_LABEL)
    old = _old_title(prs.slides[0], market)
    geo_word = old[:-len(market)].strip() if old.endswith(market) else "Global"
    geo_word = geo_word or "Global"

    if which == "tmr":
        _rebrand_text(prs)
    if which in ("tmr", "factmr"):
        _rename_shapes(prs, "TMR " if which == "tmr" else "FactMR ")
    front = prs.slides[0]
    _clear(front)
    els = _paint(front, tpl.slides[0], colours, keep_ph=(which == "factmr"))
    yrs = (config.HIST_START, config.HIST_END, config.EST_YEAR,
           config.YEARS[-1])
    if which == "tmr":
        _fill_front_tmr(els, geo_word, market, yrs, sample)
    elif which == "factmr":
        _fill_front_factmr(els, geo_word, market, yrs, sample)
    else:
        _fill_front(els, geo_word, market, yrs,
                    "Sample Report" if sample else "Report")

    # the template slides that close the deck: the contact page (FMI, TMR)
    # or the about page + contact page (Fact.MR)
    for k in ((1, 2) if which == "factmr" else (1,)):
        last = prs.slides.add_slide(prs.slide_layouts[6])
        last_els = _paint(last, tpl.slides[k], colours,
                          keep_ph=(which == "factmr"))
        if which in ("tmr", "factmr"):
            _year_in(last_els)
        for ph in list(last.placeholders):
            ph._element.getparent().remove(ph._element)

    tmp = path.with_suffix(".fl.tmp")
    prs.save(str(tmp))
    os.replace(tmp, path)
    return True
