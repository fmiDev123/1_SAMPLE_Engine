"""
Build the review workbook (FMI_Slide_Content_Templates.xlsx) from the Python templates, so the
Excel and the code never drift apart.

    python -m fmi_templates.export_workbook  [output.xlsx]  [sample_index.txt]
"""
import re
import string
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from .slide_templates import SLIDES, COMMON
from .lexicon import LEXICON
from .placeholders import describe
from .demo_muffins import render_all, ctx_for
from .rules import THRESHOLDS

NAVY = "1F3864"
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
HEAD_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Border(bottom=Side(style="thin", color="D9D9D9"))
SECTION_FILL = PatternFill("solid", fgColor="DCE6F1")

RULES_DOC = [
    ("p_trend / p_vol_trend", "F-CAGR vs H-CAGR (value / volume)", f"gap >= +{THRESHOLDS['trend_bps']} BPS: 'accelerating from'; gap <= -{THRESHOLDS['trend_bps']} BPS: 'moderating from'; otherwise 'broadly in line with'", "Replaces the hard-coded 'accelerating' on slides 4, 35, 37, 38, 41."),
    ("p_vol_period_compare / p_val_period_compare", "Forecast-decade addition vs historical addition", "ratio >= 1.1: 'exceeding'; <= 0.9: 'below'; else 'broadly matching'", "Replaces hard-coded 'more than the previous period'."),
    ("p_price_mix / p_price_mix_short", "Value F-CAGR vs volume F-CAGR", f"gap >= +{THRESHOLDS['price_mix_bps']} BPS: price and mix uplift; <= -{THRESHOLDS['price_mix_bps']}: price erosion; else stable pricing", "Used on G03, G04, M03."),
    ("p_seg1_lead_fast / p_sub_lead_fast", "Is the largest sub-segment also the fastest?", "Same: 'X both leads ... and grows fastest'; different: 'X leads ..., while Y grows fastest'", "Fixes 'Balance scale in Sweet muffins with targeted growth in Sweet muffins'."),
    ("p_seg1_fast_clause", "Same test, as a trailing clause", "Same: ', and it is also the fastest-growing category'; different: ', while Y grows fastest at ...'", "A05 row 2."),
    ("p_child_lead_fast / _short", "Is the largest child geography also the fastest?", "Same / different sentence variants", "A05, B01."),
    ("p_comm_focus / p_region_priority / p_approach / p_segment_priority", "Leader vs fastest identity", "Separate sentences for same and different cases", "G07, G08."),
    ("p_sub_share_shift / p_seg1_share_shift / p_child_share_shift_m", "Largest BPS gain", f"max gain < {THRESHOLDS['share_shift_bps']} BPS: 'mix remains broadly stable'; else 'Share shifts toward X (+n BPS) and away from Y (-n BPS)'", "One window ({est_year} to {fc_end}) on every slide."),
    ("p_item_role", "Per sub-segment or child geography in repeat lists", f"leader / fastest flags first; else size class (core >= {THRESHOLDS['core_share']:.0%}, mid >= {THRESHOLDS['niche_share']:.0%}, else niche) plus pace vs average (+/- {THRESHOLDS['role_bps']} BPS)", "Replaces 'serves a niche' / 'a smaller market' applied to every non-leader."),
    ("p_sub_landscape", "Number of sub-segments", "n <= 2: no 'top two' clause (it would always read 100%)", "D01."),
    ("p_sub_opp_tail", "Number of sub-segments", "n >= 4: names the smallest contributor; else a closing sentence", "Avoids 'adds the least' naming the same item as the top list."),
    ("p_above_pace", "Children at or above market CAGR", "0 / all / some variants", "B01, M01, M03, M07."),
    ("p_child_landscape", "Number of children", "n <= 3: no 'top three' clause", "B01."),
    ("p_vs_parent / p_bench_arrow", "Country F-CAGR vs parent region", f"+/- {THRESHOLDS['bench_bps']} BPS: outpacing / trailing / in line", "Replaces hard-coded 'trailing ... by X BPS' on slides 445-446."),
    ("p_index_leaders", "Top CAGR-index country vs top absolute-index country", "Same country: 'leads on both indices'; different: two-part sentence", "G05."),
    ("p_concentration / p_competition_line / p_k01_conc_head", "Top-3 company share (r_cr3)", f">= {THRESHOLDS['cr3_high']:.0%} highly consolidated; >= {THRESHOLDS['cr3_mid']:.0%} moderately consolidated; >= {THRESHOLDS['cr3_low']:.0%} fragmented with a leadership group; else highly fragmented", "Concentration must come from company shares, never from segment share (slide 41 bug)."),
    ("p_supply_conc", "Share of the largest child geography", f">= {THRESHOLDS['supply_conc_high']:.0%} concentrated; >= {THRESHOLDS['supply_conc_mid']:.0%} moderately concentrated; else dispersed", "M02."),
    ("p_entry_spread", "Largest child share > 50%", "Variant for one dominant geography", "M07."),
    ("p_r3_head / p_r3_body", "Concentration (child share >= 30% or seg1 lead share >= 50%)", "Concentration-risk restraint, otherwise input-cost-volatility restraint", "M04 restraint 3."),
    ("p_o3_body / p_seg2_reach", "Largest vs fastest route-to-market", "Same / different variants", "M03, M05."),
    ("p_seg2_fast_vs_market", "Fastest route F-CAGR vs market", "outpacing / in line / below", "A05 row 3."),
    ("p_price_band", "Spread between highest and lowest child price", f"< {THRESHOLDS['price_band_tight']:.0%} tight band; < {THRESHOLDS['price_band_wide']:.0%} moderate; else wide", "P05."),
    ("p_price_conv", "Do low-priced children grow price faster?", "True: converge (catch-up); False: premium widens", "P01."),
    ("p_price_dir_verb / p_price_dir_inf / p_price_outlook", "Average price CAGR", "> +0.5% rises (modest < 1%, moderate < 3%, strong); < -0.5% declines; else flat", "P02, P10."),
    ("p_price_spread_dir", "Change in high-low price gap", "+/- 5% relative: widens / narrows / holds broadly steady", "P02."),
    ("p_direction_split / p_down_list", "Factors with downward direction", "All up / some down variants", "P07, P10."),
    ("p_yoy_trend", "Average YoY, first 3 vs last 3 forecast years", "late lower by 50 BPS: moderates; higher: strengthens; else steady", "D08."),
    ("p_yoy_decline", "Sub-segments with any negative YoY", "0 / 1 / many", "D07, C05."),
    ("p_force_*", "Porter force score 0-10", "< 3.5 low; < 5.5 moderate; < 7.5 moderate to high; else high", "M07."),
    ("likelihood(score)", "Restraint or opportunity score per period", "< 4 Low; < 7 Medium; else High", "M04, M05 timelines."),
    ("p_hhi_band", "HHI", "< 1,500 fragmented; <= 2,500 moderately concentrated; else concentrated", "K11."),
    ("p_tier_word / p_t1_barrier", "Tier 1 share", ">= 50% steep; Tier 1 >= 40% 'meaningful barriers', else 'room for challengers'", "K01, L02."),
    ("p_ma_head / p_ma_body", "Players with M&A activity", ">= 3: consolidation lever; else organic growth", "K01."),
    ("p_leader_consistency", "Same leader in every share donut?", "Same / shifts variants", "L03, L04."),
    ("p_c_position", "Is the profiled company the leader?", "Leader / challenger sentence", "N02."),
    ("p_c_rev_trend", "Company revenue CAGR", "grew / declined / held broadly flat", "N04."),
    ("p_c03_takeaway / p_c04_takeaway", "Same sub-segment for two picks?", "Merged sentence when identical", "C03, C04."),
    ("p_item_drivers / p_item_restraints", "Per child geography (M01 cells)", "Assembled from flags: top-3 size, pace vs market, acceleration vs history, price growth, maturity", "Replaces keyword fragments on slide 554."),
    ("{name}_the / {name}_The", "Country takes 'the'", "U.S., U.K., UAE, Netherlands, Philippines ... (THE_NAMES in engine.py)", "'The U.S. adds ...', 'such as the Netherlands'."),
]

QC_FINDINGS = [
    ("High", "571", "Strategic activities for real companies are generated from templates (for example an acquisition by Hostess Brands in South Asia & Pacific), with 'FMI company profiling' as evidence and no date.", "K09 takes rows only from a sourced developments sheet with date and URL; otherwise prints 'No material development disclosed'."),
    ("High", "566-570", "Tier and strategy labels alternate mechanically (Buy to scale / Launch-led growth); Flowers Foods is classed as a Tier 3 local or niche player.", "Tier must come from company revenue in the input sheet; strategy from the sourced activity list or n/a."),
    ("High", "583-597", "Profile metrics are B2B sales metrics (win rate, book-to-bill, pipeline coverage, a hard-coded ~95-day sales cycle) that no secondary source reports for these companies; 'of a US$ Mn SAM' has no number; financials cite FMI Analysis and stop at FY2024.", "N01 to N05 rebuilt on sourceable data (revenue, EBITDA margin, brands, plants, countries) with c_fin_source; private-company variant for firms such as McKee Foods."),
    ("High", "583-597", "Company roster looks out of date: Hostess Brands was acquired by J.M. Smucker (completed November 2023), and ARYZTA's North American business became Aspire Bakeries in 2021, so 'ARYZTA / Aspire Bakeries' merges two separate companies.", "Roster check before each run; verify current ownership."),
    ("High", "557-558", "Real figures leak into a masked sample ('5.0%', '-52 BPS', '30.2%-share'); price axes on 536 and 540 and heatmap legends on 63/65 also print unmasked numbers.", "All market figures are v_ placeholders; qc_slide() flags any % / BPS / x number in masked mode."),
    ("High", "577, 585", "Healthcare wording in a food deck: 'hospital/clinic' channels, 'clinical validation', 'online and clinics'; 541/544 'Development-stage price assumptions' and 'Aftermarket & Service Costs'.", "Industry lexicon (lx_*) plus a banned-term list per industry in QC."),
    ("High", "4, 35, 37, 38, 41", "'accelerating' is hard-coded whatever the data; 'more than the previous period' is hard-coded.", "p_trend, p_period_compare rules."),
    ("High", "41", "'The market is consolidated' is inferred from segment share, not company concentration.", "p_competition_line uses company top-3 share, with a segment-only fallback that makes no concentration claim."),
    ("High", "234-249, 538", "Grouping headers ('Sweet Muffins Flavours / Savoury Muffins Flavours') are treated as sub-segments, producing impossible cells such as 'Sweet muffins x Savoury Muffins Flavours'.", "Cross-section columns must be leaf sub-segments; zero-by-definition cells excluded from picks (X01 rule note)."),
    ("Medium", "43", "'Fact.MR Analysis & Recommendation' heading in an FMI deck; 'Balance scale in Sweet muffins with targeted growth in Sweet muffins'.", "{publisher_short} placeholder; p_approach same/different variants."),
    ("Medium", "1", "Cover reads 'Region focus: Global' when the focus region is North America.", "{focus_region} placeholder."),
    ("Medium", "2", "Agenda group 'Country Market Outlook' and section 12 omit the country name.", "{focus_country} placeholders."),
    ("Medium", "41, 44", "'Market Overview (1/4)' and 'Scenario Outlook (4/4)' carry page tags with no sibling pages.", "{page_tag} empty unless the slide is split."),
    ("Medium", "44", "Scenario line chart titled 'Attractiveness Analysis By Region'; scenario base is 2025 while the deck base is 2026.", "G09 titles fixed to scenarios and est_year."),
    ("Medium", "445-446", "KPI reads '2025 MARKET SIZE / estimate year' while the estimate year is 2026; 'U.S.'s'.", "{est_year}; no possessive on country names; article forms."),
    ("Medium", "39, 256", "'Top 10 Countries' shown for a region with 2 countries; CAGR-index chart order differs from its table.", "{n_top} = min(10, countries); sort chart and table from one ranked list."),
    ("Medium", "35, 37, 59, 61", "Half sentences and missing words ('grows fastest XX CAGR'); window mismatch (chart 2021 to 2036, text 2026 to 2036); KPI 'CAGR 2021-2025' vs text '(2021-2026)'.", "All bullets rewritten as complete sentences; one BPS window."),
    ("Medium", "42, 554", "Key-insight boxes and dynamics cells are fragments ('2026 size XX US$ Mn; 2036 XX').", "G07 and M01 rewritten; M01 cells assembled from rules."),
    ("Medium", "46, 262", "'Sub-Segment Insights' header on region and country slides; 'Regional analysis ... across 2 countries'.", "{child_label_cap} Insights; child-aware subtitle."),
    ("Medium", "55, 541, 544, 545", "Duplicated words ('product type types', 'product type type').", "Templates fixed; QC flags duplicated words."),
    ("Medium", "559", "Rows out of chronological order (2025, 2026, 2035, 2031, 2036) and labelled 'Recent Developments' though model-derived.", "M06 renamed 'Milestones and Forecast Checkpoints', sorted; real-developments variant."),
    ("Medium", "560", "Supplier power derived from the regional share of demand.", "M07 supplier power uses input-cost lexicon and realisation trend."),
    ("Medium", "565", "Panel labelled '4. Sales Channel / By application / Who buys the output'.", "K03 seg2 panel labels fixed."),
    ("Medium", "575", "'Brand A / B / C' placeholders though brands exist on slide 10; strategy letters with no legend.", "{c_key_brands}; legend added."),
    ("Low", "447", "Bar charts padded with blank category labels when a segment has fewer than 6 sub-segments.", "Category count = sub-segment count (C02 rule note)."),
    ("Low", "452-456", "Country data slides have no narrative.", "C03 to C05 takeaways added."),
    ("Low", "51", "Stray 'By segment' labels.", "Removed; takeaway sentence added."),
    ("Low", "578", "Subtitle says 2025 while the chart is 2026; 'use the region selector' refers to interactivity a PPTX does not have.", "est_year; text points to the next slide."),
    ("Low", "8 vs 41", "'17 axes' on the taxonomy slide vs '21 segmentation lenses' in the overview.", "One n_segments value used everywhere."),
    ("Low", "All", "En dashes in titles and ranges (2021-2036 written with an en dash; 'Industry Structure Landscape -').", "Hyphens only; clean() replaces any en or em dash."),
]


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")


def _sheet(wb, title, headers, rows, widths, freeze="A2"):
    ws = wb.create_sheet(title)
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font, cell.alignment, cell.border = BODY_FONT, WRAP, THIN
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions
    return ws


def _flatten(val):
    """Yield (index, text, is_repeat) for an element value."""
    if isinstance(val, dict) and "repeat" in val:
        yield 1, f"[repeat for each item in {val['repeat']}] {val['text']}", True
    elif isinstance(val, dict):
        for i, (k, v) in enumerate(val.items(), 1):
            yield i, f"{k}: {v}", False
    elif isinstance(val, (list, tuple)):
        for i, v in enumerate(val, 1):
            if isinstance(v, dict) and "repeat" in v:
                yield i, f"[repeat for each item in {v['repeat']}] {v['text']}", True
            elif isinstance(v, (list, tuple)):
                yield i, " | ".join(map(str, v)), False
            else:
                yield i, v, False
    else:
        yield 1, val, False


SLIDE_MAP_PATTERNS = [
    (r"^Section \d", "A03"), (r"^Global .*Market$", "A01"), (r"^Agenda$", "A02"), (r"^Market Signals", "A05"), (r"^Market Taxonomy Regions", "A06"),
    (r"^Market Taxonomy Segments", "A07"), (r"Company Coverage", "A08"), (r"^Table of Contents", "A09"),
    (r"^List of Figures", "A10"), (r"^List of Tables", "A11"),
    (r"Market Summary, By Volume", "G01"), (r"Market Pricing & Regional Price Index", "G02"), (r"Market Summary, By Value", "G03"),
    (r"Market Volume & Value Analysis", "G04"), (r"Country Indices", "G05"), (r"^Market Overview", "G06"),
    (r"^Supply Side Trends", "G07"), (r"^Analysis and Recommendation", "G08"), (r"^Scenario Outlook", "G09"),
    (r"(Regional|Country) Outlook$", "B01"), (r"Market Analysis, By (Region|Country)", "B02"),
    (r"Market Share Analysis, By (Region|Country)", "B03"), (r"Market Trends, By (Region|Country)", "B04"), (r"Top \d+ Countries", "B05"),
    (r"Segment Outlook", "D01"), (r"Market Structure & Attractiveness", "D02"), (r"Absolute \$ Opportunity Analysis", "D03"),
    (r"Market Volume Analysis, By", "D04"), (r"Market Value Analysis, By", "D05"), (r"Value & Share Shift", "D06"),
    (r"^YoY Analysis of .* \(%\) \(2/2\)", "D08"), (r"^YoY Analysis of", "D07"),
    (r"Market Snapshot: ", "X01"), (r"Market Share Analysis: ", "X02"), (r"Absolute \$ Opportunity & Attractiveness", "X03"),
    (r"Volume Overview|Value Overview", "C01"), (r"Key Segments Overview", "C02"), (r"Segmental Analysis", "C03"),
    (r"Market Tables", "C04"), (r"YoY Analysis [A-Z]", "C05"),
    (r"Product Pricing & Regional Price Index Trends", "P01"), (r"Pricing Assessment: By", "P02"), (r"Price Matrix", "P03"),
    (r"Region Benchmarking", "P04"), (r"Price Positioning & Key Findings", "P05"), (r"Key Pricing Drivers", "P06"),
    (r"Factors Impacting Pricing", "P07"), (r"Regional Price Assessment", "P08"), (r"Global Pricing Benchmarking", "P09"),
    (r"Price Forecast & Dynamics", "P10"), (r"Competitor Pricing Strategies", "P11"),
    (r"Value Chain & Stakeholder", "V01"), (r"Operating Margin", "V02"), (r"Buyers Assessment", "V03"),
    (r"Dynamics by Region", "M01"), (r"Driving Factors \(1/2\)", "M02"), (r"Driving Factors \(2/2\)", "M03"),
    (r"Restraints$", "M04"), (r"Opportunities$", "M05"), (r"Recent Developments", "M06"), (r"Porter", "M07"),
    (r"at a Glance", "K01"), (r"Segment Leadership Map", "K02"), (r"Value Chain Positioning", "K03"), (r"Player Landscape", "K04"),
    (r"Corporate Scorecard", "K05"), (r"Offering Scorecard", "K06"), (r"Execution Scorecard", "K07"), (r"Success Factors", "K08"),
    (r"Strategic Activity Tracker", "K09"), (r"Strategic Activity by Function", "K10"), (r"Industry Structure Landscape .*\(1/2\)", "K11"),
    (r"Industry Structure Landscape .*\(2/2\)", "K12"), (r"Competitive Dashboard", "K13"),
    (r"Advice for Market Vendors", "L01"), (r"Company Value Share", "L02"), (r"Company Share by Product", "L03"),
    (r"Company Share by Region", "L04"), (r"Who Shaped Growth", "L05"),
    (r"Executive Market Snapshot", "N01"), (r"Regional Penetration", "N02"), (r"SWOT", "N03"), (r"Key Financials", "N04"),
    (r"Strategy Analysis", "N05"), (r"Market$", "A04"),
]


def slide_map(index_path):
    rows = []
    for line in open(index_path, encoding="utf-8"):
        parts = [p.strip() for p in line.split("|")]
        if len(parts) < 2:
            continue
        n, title = parts[0], parts[1]
        tid = next((t for pat, t in SLIDE_MAP_PATTERNS if re.search(pat, title)), "?")
        rows.append((int(n), title, tid))
    return rows


def build(out_path, index_path=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    by_section = {}
    for s in SLIDES:
        by_section.setdefault(s["section"], []).append(s["id"])
    readme = [
        ("FMI Auto-Deck Slide Content Templates", None),
        ("What this is", "Industry-neutral content for every slide type in the FMI Python-generated report deck (76 types cover all 597 slides of the Packaged Muffins sample). Text is written once with placeholders; wording that depends on the data is chosen by rules, and industry terms come from a lexicon, so the deck needs no AI or manual correction."),
        ("Source of truth", "The Python package fmi_templates (slide_templates.py, rules.py, lexicon.py, engine.py). This workbook is generated from it for review; edit the .py files, then re-export."),
        ("How the pipeline uses it", "1) Build a flat context dict per slide (names, n_ counts, v_ formatted figures, r_ raw numbers, c_ company data). 2) ctx = build_ctx(ctx) adds lexicon, article forms and rule phrases. 3) render_slide(id, ctx, mask=True/False) returns element -> text. 4) qc_slide() must return no issues before a deck is saved."),
        ("Placeholder prefixes", "v_ market figure (masked as XX in sample decks) | c_ real company data (never masked) | n_ count | s_ model score | p_ rule phrase | lx_ industry wording | r_ raw number for rules (never printed)."),
        ("Sample vs full deck", "mask=True prints every v_ value as XX (sample reports). mask=False requires every v_ value, and rendering fails loudly if one is missing."),
        ("Style rules built in", "No em or en dashes; every bullet is a complete sentence; ranges written as 2026-2036; one BPS window (est_year to fc_end); no possessives on country names; 'the' added automatically for the U.S., U.K., Netherlands and similar."),
        ("Sheets", "Templates: every element of every slide type. Rendered Example: the muffins sample re-rendered (masked). Slide Map: all 597 sample slides mapped to a template ID. Rules: the word-choice logic. Placeholders: definitions. Lexicon: industry wording for food, consumer, healthcare, pharma and industrial. Sample QC Findings: defects in the current sample and how the templates fix them."),
        ("Inputs still needed from secondary research", "Company profiles (revenue, EBITDA margin, plants, countries, brands, stated strategy, strengths and weaknesses) and the developments list (date, company, development, source URL). The templates leave bracketed slots for these and never invent them."),
        ("Slide types by section", "; ".join(f"{k}: {', '.join(v)}" for k, v in by_section.items())),
    ]
    for i, (k, v) in enumerate(readme, 1):
        ws.cell(row=i, column=1, value=k).font = Font(name="Arial", bold=True, size=14 if i == 1 else 10, color=NAVY if i == 1 else "000000")
        if v:
            c = ws.cell(row=i, column=2, value=v)
            c.font, c.alignment = BODY_FONT, WRAP
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 120

    # Templates
    rows = []
    for s in SLIDES:
        base = (s["id"], s["section"], s["name"], s["applies"], s["sample"])
        rows.append(base + ("title", 1, s["title"]))
        if s["subtitle"]:
            rows.append(base + ("subtitle", 1, s["subtitle"]))
        for key, val in s["elements"].items():
            for i, text, _ in _flatten(val):
                rows.append(base + (key, i, text))
    rows.append(("ALL", "Common", "Every slide", "all", "all", "geo_tag", 1, COMMON["geo_tag"]))
    rows.append(("ALL", "Common", "Every slide", "all", "all", "source", 1, COMMON["source"]))
    rows.append(("ALL", "Common", "Every slide", "all", "all", "footer", 1, COMMON["footer"]))
    wst = _sheet(wb, "Templates", ["Template ID", "Section", "Slide type", "Applies to", "Sample slides", "Element", "#", "Template text"],
                 rows, [10, 16, 30, 14, 16, 18, 5, 110], "C2")
    last = None
    for r in range(2, wst.max_row + 1):
        tid = wst.cell(row=r, column=1).value
        if tid != last:
            for c in range(1, 9):
                wst.cell(row=r, column=c).fill = SECTION_FILL
            wst.cell(row=r, column=8).font = BOLD
            last = tid

    # Rendered example
    rrows = []
    for r, _ in render_all(mask=True):
        for k, v in r.items():
            if k in ("id", "name", "_missing", "geo_tag", "source", "footer"):
                continue
            for i, text, _ in _flatten(v if not isinstance(v, dict) else v):
                rrows.append((r["id"], r["name"], k, i, text))
    _sheet(wb, "Rendered Example", ["Template ID", "Slide type", "Element", "#", "Rendered text (Packaged Muffins, masked sample, food lexicon)"],
           rrows, [10, 30, 18, 5, 120], "C2")

    # Slide map
    if index_path:
        names = {s["id"]: s["name"] for s in SLIDES}
        smap = [(n, t, tid, names.get(tid, "")) for n, t, tid in slide_map(index_path)]
        _sheet(wb, "Slide Map", ["Sample slide #", "Sample slide title", "Template ID", "Template slide type"], smap, [12, 80, 12, 40])

    _sheet(wb, "Rules", ["Phrase key", "Decides on", "Logic", "Where / why"], RULES_DOC, [34, 34, 70, 50])

    # Placeholders
    import string as _s
    F = _s.Formatter()
    used = {}

    def walk(v, sid):
        if isinstance(v, str):
            for _, f, _, _ in F.parse(v):
                if f:
                    used.setdefault(f, set()).add(sid)
        elif isinstance(v, dict):
            for k, x in v.items():
                walk(k, sid)
                walk(x, sid)
        elif isinstance(v, (list, tuple)):
            for x in v:
                walk(x, sid)

    for s in SLIDES:
        walk(s["title"], s["id"])
        walk(s["subtitle"], s["id"])
        walk(s["elements"], s["id"])
    demo = ctx_for("G06")
    prow = []
    for name in sorted(used, key=lambda x: (x.split("_")[0] not in ("v", "c", "n", "s", "p", "lx", "r"), x)):
        kind, desc = describe(name)
        ex = demo.get(name, "")
        if name.startswith("v_"):
            ex = "XX in samples"
        prow.append((name, kind, desc, str(ex)[:120], ", ".join(sorted(used[name]))[:200]))
    _sheet(wb, "Placeholders", ["Placeholder", "Type", "Definition", "Example (muffins)", "Used in templates"], prow, [30, 12, 70, 40, 50])

    # Lexicon
    keys = [k for k in LEXICON["food"] if not isinstance(LEXICON["food"][k], list)]
    lrows = [(k,) + tuple(LEXICON[ind][k] for ind in LEXICON) for k in keys]
    for k in ("lx_buyer_types", "lx_buyer_factors"):
        lrows.append((k,) + tuple("; ".join(LEXICON[ind][k]) for ind in LEXICON))
    for i in range(5):
        lrows.append((f"lx_vc{i + 1} (stage | description)",) + tuple(f"{LEXICON[ind]['lx_vc'][i][0]} | {LEXICON[ind]['lx_vc'][i][1]}" for ind in LEXICON))
    _sheet(wb, "Lexicon", ["Key"] + list(LEXICON), lrows, [26] + [44] * len(LEXICON), "B2")

    _sheet(wb, "Sample QC Findings", ["Severity", "Sample slides", "Issue found in the current sample", "How the templates fix it"],
           QC_FINDINGS, [10, 16, 80, 70])

    wb.save(out_path)
    return out_path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "FMI_Slide_Content_Templates.xlsx"
    idx = sys.argv[2] if len(sys.argv) > 2 else None
    print(build(out, idx))
