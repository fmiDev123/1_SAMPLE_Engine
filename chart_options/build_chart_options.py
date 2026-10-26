"""
Chart-options sampler  -  standalone, touches nothing in the main project.

Reads the Performance Knitted Fabrics formatted dataset and writes ONE short
deck (one slide per chart type) showing which visual can replace / upgrade the
charts the report uses today.  Native PowerPoint charts wherever possible
(editable in PowerPoint); shape-built where PowerPoint has no native type
(treemap, marimekko, dumbbell, heat-map, tile map).

Run:   python chart_options/build_chart_options.py
Out:   chart_options/Chart_Options_Performance_Knitted_Fabrics.pptx
"""
from __future__ import annotations

import copy
import sys
from pathlib import Path

import pandas as pd
from lxml import etree
from pptx import Presentation
from pptx.chart.data import BubbleChartData, CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DATA = (ROOT / "Input" / "Performance_Knitted_Fabrics_Market" /
        "Global_Performance_Knitted_Fabrics_Market_Formatted_Dataset.xlsx")
OUT = Path(__file__).resolve().parent / "Chart_Options_Performance_Knitted_Fabrics.pptx"

# ---- palette (same family as fmi_report/theme_fmi.py) ---------------------
INK, BLUE, TEAL = "0B2E4A", "0E72AE", "008189"
OCHRE, RUST, PURPLE = "A86F00", "A35313", "6F3CC2"
SLATE, CORAL, SKY = "6C7A89", "E44D3A", "3D85C6"
GREY, LIGHT, PANEL = "B2B2B2", "E6EBF0", "F5F8FA"
WHITE, RED = "FFFFFF", "C00000"
CAT = [BLUE, TEAL, OCHRE, PURPLE, RUST, SLATE, SKY, CORAL, INK]
YEARS = [str(y) for y in range(2021, 2037)]


def rgb(h): return RGBColor.from_string(h)


# ---------------------------------------------------------------- data ----
def load():
    d = pd.read_excel(DATA, sheet_name="Master Sheet", header=5)
    d.columns = [str(c) for c in d.columns]
    d = d.dropna(subset=["Region"])
    return d


D = load()


def series(region, segment, unit="Value"):
    """parent-level rows -> DataFrame index=Sub-segment, columns=years."""
    s = D[(D.Region == region) & (D.Segment == segment) & (D.Unit == unit)]
    return s.groupby("Sub-segment", sort=False)[YEARS].sum()


def total(region="Global", unit="Value"):
    return series(region, "Region" if region == "Global" else "Material", unit).sum()


def cagr(a, b, n): return (b / a) ** (1 / n) - 1 if a > 0 else 0.0


REGIONS = series("Global", "Region")
MATERIAL = series("Global", "Material")
ENDUSE = series("Global", "End-Use Application")
FUNCTION = series("Global", "Primary Performance Function")
CONSTR = series("Global", "Fabric Construction")
WEIGHT = series("Global", "Fabric Weight")
GTOT = REGIONS.sum()
VOL_REG = series("Global", "Region", "Volume")
USxM = D[(D.Region == "Global") & (D.Segment == "End-Use Application Vs Material")
         & (D.Unit == "Value")]

COUNTRIES = D[(D.Segment == "Country") & (D.Unit == "Value")
              & (D.Region == "Global")].set_index("Sub-segment")[YEARS]


def short(name, n=22):
    return name if len(name) <= n else name[: n - 1].rstrip() + "…"


# ------------------------------------------------------------- helpers ----
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
CH_L, CH_T, CH_W, CH_H = Inches(0.45), Inches(1.45), Inches(8.55), Inches(5.55)
PN_L, PN_W = Inches(9.3), Inches(3.6)


def tb(slide, x, y, w, h, text, size=12, bold=False, color=INK, align=PP_ALIGN.LEFT,
       anchor=MSO_ANCHOR.TOP, italic=False, wrap=True):
    s = slide.shapes.add_textbox(x, y, w, h)
    tf = s.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold, r.font.italic = Pt(size), bold, italic
        r.font.color.rgb = rgb(color)
        r.font.name = "Calibri"
    return s


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    return s


def box_text(s, text, size=10, color=WHITE, bold=False, align=PP_ALIGN.CENTER):
    tf = s.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.03)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = ln
        r.font.size, r.font.bold = Pt(size), bold
        r.font.color.rgb = rgb(color)
        r.font.name = "Calibri"


def line(slide, x1, y1, x2, y2, color=GREY, w=1.0, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.line.color.rgb = rgb(color)
    c.line.width = Pt(w)
    if dash:
        c.line.dash_style = 4
    return c


def new_slide(title, subtitle, kind, best_for, why, instead, tag=None):
    s = prs.slides.add_slide(BLANK)
    tb(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.6), title, 24, True)
    tb(s, Inches(0.45), Inches(0.88), Inches(12.4), Inches(0.4), subtitle, 13, color=SLATE)
    line(s, Inches(0.45), Inches(1.33), Inches(12.88), Inches(1.33), BLUE, 1.5)
    # right-hand "how to choose" panel
    p = box(s, PN_L, CH_T, PN_W, CH_H, PANEL, LIGHT)
    y = CH_T + Inches(0.12)
    for head, body, col in (("CHART TYPE", kind, BLUE), ("BEST FOR", best_for, TEAL),
                            ("WHY IT FITS THIS DATA", why, OCHRE),
                            ("USE IT INSTEAD OF", instead, RUST)):
        tb(s, PN_L + Inches(0.12), y, PN_W - Inches(0.24), Inches(0.25), head, 9.5, True, col)
        h = Inches(0.28 + 0.215 * max(1, len(body) // 38 + 1))
        tb(s, PN_L + Inches(0.12), y + Inches(0.24), PN_W - Inches(0.24), h, body, 11.5)
        y += Inches(0.3) + h
    if tag:
        tb(s, Inches(0.45), Inches(7.08), Inches(8.5), Inches(0.3), tag, 9, color=SLATE, italic=True)
    return s


def style_chart(ch, legend=True, legend_pos=XL_LEGEND_POSITION.BOTTOM, size=10):
    ch.font.size = Pt(size)
    ch.font.name = "Calibri"
    ch.font.color.rgb = rgb(INK)
    ch.has_legend = legend
    if legend:
        ch.legend.position = legend_pos
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(size)


def clean_axes(ch, val_visible=True, grid=True, num_fmt='#,##0'):
    try:
        ca, va = ch.category_axis, ch.value_axis
    except Exception:
        return
    ca.format.line.color.rgb = rgb(GREY)
    ca.has_major_gridlines = False
    ca.tick_labels.font.size = Pt(10)
    va.has_major_gridlines = grid
    if grid:
        va.major_gridlines.format.line.color.rgb = rgb(LIGHT)
    va.format.line.fill.background()
    va.tick_labels.number_format = num_fmt
    va.tick_labels.number_format_is_linked = False
    va.tick_labels.font.size = Pt(10)
    va.visible = val_visible


def colour_series(ch, colours):
    for s, c in zip(ch.plots[0].series, colours):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)


def labels(plot, fmt='#,##0', size=9, pos=None, color=INK, bold=False):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size, dl.font.color.rgb, dl.font.bold = Pt(size), rgb(color), bold
    if pos is not None:
        dl.position = pos


def add_chart(slide, ctype, cd, x=CH_L, y=CH_T, w=CH_W, h=CH_H):
    return slide.shapes.add_chart(ctype, x, y, w, h, cd).chart


def ramp(v, lo, hi, c0=(234, 246, 251), c1=(11, 46, 74)):
    t = 0 if hi == lo else max(0.0, min(1.0, (v - lo) / (hi - lo)))
    return "%02X%02X%02X" % tuple(int(a + (b - a) * t) for a, b in zip(c0, c1)), t


# ============================================================ SLIDES ======
# ---- 0. cover + option matrix ---------------------------------------------
s = prs.slides.add_slide(BLANK)
box(s, 0, 0, prs.slide_width, prs.slide_height, INK)
tb(s, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.2),
   "Chart options for the report & sample", 40, True, WHITE)
tb(s, Inches(0.8), Inches(3.7), Inches(11.5), Inches(1.0),
   ["One example per chart type, built on the Global Performance Knitted Fabrics dataset",
    "What can replace or upgrade the bar / stacked / line charts we use today"], 16, False, "BFD8EA")
tb(s, Inches(0.8), Inches(6.6), Inches(11.5), Inches(0.4),
   "Concept sheet - separate from the report engine. Figures are live from the formatted dataset.",
   11, False, "8FB0C8", italic=True)

s = prs.slides.add_slide(BLANK)
tb(s, Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.6), "Upgrade map: what we use today → what can replace it", 24, True)
line(s, Inches(0.45), Inches(0.95), Inches(12.88), Inches(0.95), BLUE, 1.5)
rows = [
    ("Question the client asks", "Used today", "Better option(s)", "Slide"),
    ("Which region / segment is biggest?", "Clustered column", "Ranked horizontal bar with CAGR tag", "3"),
    ("How does the mix change over time?", "Stacked column", "100% stacked column; stacked area", "4, 6"),
    ("How fast is it growing?", "Line", "Line + forecast shading + end labels; Combo (value + YoY)", "5, 7"),
    ("How big is each share?", "Pie / doughnut", "Doughnut with centre KPI; Treemap", "8, 10"),
    ("What drives the growth?", "-", "Waterfall (growth bridge)", "9"),
    ("Where are size and growth together?", "Table", "Bubble chart (size × CAGR × value)", "11"),
    ("Which cell of a matrix matters?", "Table", "Heat-map; Marimekko", "12, 13"),
    ("Who moved up / down?", "Table", "Slope chart; Dumbbell", "14, 15"),
    ("Profile / comparison", "-", "Radar; Small multiples", "16, 17"),
    ("Geography at a glance", "Bar by region", "Tile map", "18"),
    ("Headline numbers", "Text boxes", "KPI cards with sparkline", "19"),
]
tbl = s.shapes.add_table(len(rows), 4, Inches(0.45), Inches(1.15), Inches(12.43), Inches(5.6)).table
for j, w in enumerate((3.7, 2.3, 5.6, 0.83)):
    tbl.columns[j].width = Inches(w)
for i, r in enumerate(rows):
    for j, t in enumerate(r):
        c = tbl.cell(i, j)
        c.text = t
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(INK if i == 0 else (PANEL if i % 2 else WHITE))
        p = c.text_frame.paragraphs[0]
        p.runs[0].font.size = Pt(13)
        p.runs[0].font.bold = (i == 0 or j == 2)
        p.runs[0].font.color.rgb = rgb(WHITE if i == 0 else INK)
        p.runs[0].font.name = "Calibri"
        c.vertical_anchor = MSO_ANCHOR.MIDDLE

# ---- 1/2 bar: today vs upgrade (two slides side by side on one) -----------
reg26 = REGIONS["2026"].sort_values()
reg36 = REGIONS["2036"]
s = new_slide("Bar chart - from plain columns to a ranked, labelled bar",
              "Global performance knitted fabrics value by region, 2026 (US$ Mn)",
              "Ranked horizontal bar + highlight + CAGR tag",
              "Comparing sizes across 5-15 categories with long names",
              "Region and segment names are long; sorted bars read as a league table and the lead item pops.",
              "Unsorted clustered column / vertical bar",
              "Left: today's style.  Right: upgraded.")
cd = CategoryChartData()
cd.categories = [short(c, 14) for c in REGIONS.index]
cd.add_series("2026", [round(v) for v in REGIONS["2026"]])
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd, CH_L, CH_T, Inches(3.6), CH_H)
style_chart(ch, False, size=8)
clean_axes(ch)
colour_series(ch, [GREY])
labels(ch.plots[0], size=8)
ch.has_title = True
ch.chart_title.text_frame.text = "Today"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)
cd = CategoryChartData()
order = reg26.index.tolist()
cd.categories = [f"{short(r, 22)}   ({cagr(REGIONS.loc[r,'2026'], REGIONS.loc[r,'2036'], 10)*100:.1f}% CAGR)" for r in order]
cd.add_series("2026", [round(v) for v in reg26])
ch = add_chart(s, XL_CHART_TYPE.BAR_CLUSTERED, cd, CH_L + Inches(3.7), CH_T, Inches(4.85), CH_H)
style_chart(ch, False)
clean_axes(ch, val_visible=False, grid=False)
colour_series(ch, [SKY])
pts = ch.plots[0].series[0].points
for i in range(len(order)):
    pts[i].format.fill.solid()
    pts[i].format.fill.fore_color.rgb = rgb(INK if i >= len(order) - 2 else SKY)
ch.plots[0].gap_width = 45
labels(ch.plots[0], size=10, pos=XL_LABEL_POSITION.OUTSIDE_END, bold=True)
ch.has_title = True
ch.chart_title.text_frame.text = "Upgraded: sorted, top-2 highlighted, CAGR in label"
ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

# ---- 100% stacked + stacked column ---------------------------------------
s = new_slide("Stacked column → 100% stacked (share shift)",
              "Global value by material, 2021 / 2026 / 2031 / 2036 - share of market (%)",
              "100% stacked column with in-bar % labels",
              "Showing how the MIX changes when totals differ a lot",
              "Absolute stacks hide the small segments; 100% bars make Polyester vs Nylon vs Elastane share moves obvious.",
              "Stacked column of absolute value for mix questions",
              "Pair with the absolute stacked chart when size also matters.")
parent = MATERIAL.groupby(level=0, sort=False).sum()
yrs = ["2021", "2026", "2031", "2036"]
cd = CategoryChartData()
cd.categories = yrs
for m in parent.index:
    cd.add_series(short(m, 24), [round(parent.loc[m, y] / parent[y].sum() * 100, 1) for y in yrs])
ch = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED_100, cd)
style_chart(ch, True, XL_LEGEND_POSITION.RIGHT)
clean_axes(ch, num_fmt='0"%"')
colour_series(ch, CAT)
ch.plots[0].gap_width = 55
ch.plots[0].overlap = 100
labels(ch.plots[0], '0.0"%"', 9, color=WHITE)

# ---- line: upgraded -------------------------------------------------------
s = new_slide("Line chart - add forecast shading and end labels",
              "Global value by region, 2021-2036 (US$ Mn)",
              "Multi-line with forecast zone, direct end labels",
              "Trend and ranking over 10+ periods",
              "2021-2025 is actual and 2026-2036 forecast; shading says so without a footnote.",
              "Plain line chart with a legend only",
              "End labels remove eye travel to the legend.")
top = REGIONS.sort_values("2036", ascending=False).index[:5]
cd = CategoryChartData()
cd.categories = YEARS
for r in top:
    cd.add_series(r, [round(v) for v in REGIONS.loc[r, YEARS]])
ch = add_chart(s, XL_CHART_TYPE.LINE, cd)
style_chart(ch, False)
clean_axes(ch)
for i, ser in enumerate(ch.plots[0].series):
    ser.format.line.color.rgb = rgb(CAT[i])
    ser.format.line.width = Pt(2.5)
    ser.smooth = False
    last = ser.points[len(YEARS) - 1]
    last.data_label.has_text_frame = True
    last.data_label.text_frame.text = f"{short(top[i], 16)}  {REGIONS.loc[top[i], '2036']:,.0f}"
    last.data_label.position = XL_LABEL_POSITION.RIGHT
    last.data_label.text_frame.paragraphs[0].runs[0].font.size = Pt(9)
    last.data_label.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(CAT[i])
    last.data_label.text_frame.paragraphs[0].runs[0].font.bold = True
ch.value_axis.minimum_scale = 0
# forecast band is a plain shape placed behind the plot area (editable)
band = box(s, CH_L + Inches(0.78) + int((CH_W - Inches(2.9)) * 5 / 16), CH_T + Inches(0.12),
           int((CH_W - Inches(2.9)) * 11 / 16), CH_H - Inches(0.7), "EAF6FB")
s.shapes._spTree.remove(band._element)
s.shapes._spTree.insert(2, band._element)
tb(s, CH_L + Inches(3.6), CH_T + Inches(0.1), Inches(2.5), Inches(0.3), "FORECAST PERIOD", 9, True, BLUE)

# ---- stacked area ----------------------------------------------------------
s = new_slide("Area chart - cumulative market build-up",
              "Global value by end-use application, 2021-2036 (US$ Mn)",
              "Stacked area",
              "Total market size AND contribution of each part over time",
              "Seven end-uses stack into one silhouette: the top edge is the market, bands show who feeds the growth.",
              "Stacked column when there are 12+ periods",
              "Keep to 5-7 bands; group the rest as 'Other'.")
eu = ENDUSE.groupby(level=0, sort=False).sum()
cd = CategoryChartData()
cd.categories = YEARS
for m in eu.index:
    cd.add_series(short(m, 28), [round(v) for v in eu.loc[m, YEARS]])
ch = add_chart(s, XL_CHART_TYPE.AREA_STACKED, cd)
style_chart(ch, True, XL_LEGEND_POSITION.BOTTOM)
clean_axes(ch)
for ser, c in zip(ch.plots[0].series, CAT):
    ser.format.fill.solid()
    ser.format.fill.fore_color.rgb = rgb(c)

# ---- combo: value column + YoY line (XML) ---------------------------------
s = new_slide("Combo - value columns with YoY growth line",
              "Global performance knitted fabrics, value (US$ Mn) and YoY growth (%)",
              "Column + line on a secondary axis",
              "Size AND momentum in a single frame",
              "Value climbs every year while the YoY line shows growth accelerating from about 6% to 7.5% - two messages, one picture.",
              "Two separate charts (value, then growth table)",
              "Column = left axis, line = right axis.")
tot = GTOT
yoy = [None] + [round((tot[YEARS[i]] / tot[YEARS[i - 1]] - 1) * 100, 1) for i in range(1, 16)]
cd = CategoryChartData()
cd.categories = YEARS
cd.add_series("Value (US$ Mn)", [round(v) for v in tot[YEARS]])
cd.add_series("YoY growth (%)", yoy)
ch = add_chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, cd)
style_chart(ch, True, XL_LEGEND_POSITION.BOTTOM)
clean_axes(ch)
ch.plots[0].series[0].format.fill.solid()
ch.plots[0].series[0].format.fill.fore_color.rgb = rgb(SKY)
ch.plots[0].gap_width = 60
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
ns = {"c": C}
plot = ch._chartSpace.find(".//c:plotArea", ns)
bar = plot.find("c:barChart", ns)
sers = bar.findall("c:ser", ns)
ser2 = sers[1]
bar.remove(ser2)
ax_ids = [a.get("val") for a in bar.findall("c:axId", ns)]
lc = etree.SubElement(plot, f"{{{C}}}lineChart")
etree.SubElement(lc, f"{{{C}}}grouping").set("val", "standard")
etree.SubElement(lc, f"{{{C}}}varyColors").set("val", "0")
# series colour / marker
for el in ser2.findall("c:spPr", ns):
    ser2.remove(el)
sp = etree.fromstring(
    f'<c:spPr xmlns:c="{C}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    f'<a:ln w="31750"><a:solidFill><a:srgbClr val="{CORAL}"/></a:solidFill></a:ln></c:spPr>')
ser2.insert(list(ser2).index(ser2.find("c:tx", ns)) + 1, sp)
for tag in ("invertIfNegative",):
    for el in ser2.findall(f"c:{tag}", ns):
        ser2.remove(el)
mk = etree.fromstring(
    f'<c:marker xmlns:c="{C}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
    f'<c:symbol val="circle"/><c:size val="6"/><c:spPr><a:solidFill><a:srgbClr val="{CORAL}"/></a:solidFill>'
    f'</c:spPr></c:marker>')
ser2.insert(list(ser2).index(sp) + 1, mk)
lc.append(ser2)
etree.SubElement(lc, f"{{{C}}}marker").set("val", "1")
etree.SubElement(lc, f"{{{C}}}axId").set("val", "50010")
etree.SubElement(lc, f"{{{C}}}axId").set("val", "50020")
# place lineChart right after barChart, then secondary axes at the end
plot.remove(lc)
plot.insert(list(plot).index(bar) + 1, lc)
cat2 = etree.fromstring(
    f'<c:catAx xmlns:c="{C}"><c:axId val="50010"/><c:scaling><c:orientation val="minMax"/></c:scaling>'
    f'<c:delete val="1"/><c:axPos val="b"/><c:majorTickMark val="none"/><c:minorTickMark val="none"/>'
    f'<c:tickLblPos val="nextTo"/><c:crossAx val="50020"/><c:crosses val="autoZero"/>'
    f'<c:auto val="1"/><c:lblAlgn val="ctr"/><c:lblOffset val="100"/><c:noMultiLvlLbl val="0"/></c:catAx>')
val2 = etree.fromstring(
    f'<c:valAx xmlns:c="{C}"><c:axId val="50020"/><c:scaling><c:orientation val="minMax"/>'
    f'<c:min val="0"/></c:scaling><c:delete val="0"/><c:axPos val="r"/>'
    f'<c:numFmt formatCode="0&quot;%&quot;" sourceLinked="0"/><c:majorTickMark val="none"/>'
    f'<c:minorTickMark val="none"/><c:tickLblPos val="nextTo"/><c:crossAx val="50010"/>'
    f'<c:crosses val="max"/><c:crossBetween val="between"/></c:valAx>')
last_ax = [e for e in plot if e.tag.endswith("Ax")][-1]
idx = list(plot).index(last_ax) + 1
plot.insert(idx, cat2)
plot.insert(idx + 1, val2)

# ---- doughnut with centre KPI ---------------------------------------------
s = new_slide("Pie → doughnut with the total in the centre",
              "Global value by fabric construction & weight, 2026 (US$ Mn)",
              "Doughnut + centre KPI",
              "Part-to-whole with 2-6 slices",
              "Warp vs weft and three weight classes are small sets: the share is the story and the total sits in the hole.",
              "Pie chart / 3-D pie",
              "Over 6 slices → use bar or treemap.")
cd = CategoryChartData()
cd.categories = list(CONSTR.index)
cd.add_series("2026", [round(v) for v in CONSTR["2026"]])
ch = add_chart(s, XL_CHART_TYPE.DOUGHNUT, cd, CH_L, CH_T, Inches(4.2), CH_H)
style_chart(ch, True, XL_LEGEND_POSITION.BOTTOM)
for i, c in enumerate((INK, SKY)):
    p = ch.plots[0].series[0].points[i]
    p.format.fill.solid()
    p.format.fill.fore_color.rgb = rgb(c)
labels(ch.plots[0], '0%', 12, color=WHITE, bold=True)
ch.plots[0].data_labels.show_percentage = True
ch.plots[0].data_labels.show_value = False
tb(s, CH_L + Inches(1.05), CH_T + Inches(2.05), Inches(2.1), Inches(0.9),
   [f"US$ {CONSTR['2026'].sum():,.0f} Mn", "2026"], 15, True, INK, PP_ALIGN.CENTER)
cd = CategoryChartData()
cd.categories = [short(c, 24) for c in WEIGHT.index]
cd.add_series("2026", [round(v) for v in WEIGHT["2026"]])
ch = add_chart(s, XL_CHART_TYPE.DOUGHNUT, cd, CH_L + Inches(4.3), CH_T, Inches(4.2), CH_H)
style_chart(ch, True, XL_LEGEND_POSITION.BOTTOM)
for i, c in enumerate((TEAL, OCHRE, RUST)):
    p = ch.plots[0].series[0].points[i]
    p.format.fill.solid()
    p.format.fill.fore_color.rgb = rgb(c)
labels(ch.plots[0], '0%', 12, color=WHITE, bold=True)
ch.plots[0].data_labels.show_percentage = True
ch.plots[0].data_labels.show_value = False
tb(s, CH_L + Inches(5.35), CH_T + Inches(2.05), Inches(2.1), Inches(0.9),
   [f"US$ {WEIGHT['2026'].sum():,.0f} Mn", "2026"], 15, True, INK, PP_ALIGN.CENTER)

# ---- waterfall -------------------------------------------------------------
s = new_slide("Waterfall - where the growth comes from",
              "Global value bridge 2026 → 2036 by material (US$ Mn)",
              "Waterfall / growth bridge (stacked column, hidden base)",
              "Explaining the change between two years",
              "Answers 'who adds the US$ Mn?': Polyester and Nylon contribute most of the 2026-2036 increase.",
              "Two separate bars + a table of differences",
              "Great for the 'Opportunity' / absolute-$ slides.")
mats = parent.index.tolist()
delta = [parent.loc[m, "2036"] - parent.loc[m, "2026"] for m in mats]
order = sorted(range(len(mats)), key=lambda i: -delta[i])
cats = ["2026"] + [short(mats[i], 18) for i in order] + ["2036"]
base, tot_s, inc = [None], [round(parent["2026"].sum())], [None]
run = parent["2026"].sum()
for i in order:
    base.append(round(run))
    inc.append(round(delta[i]))
    tot_s.append(None)
    run += delta[i]
base.append(None)
tot_s.append(round(run))
inc.append(None)
cd = CategoryChartData()
cd.categories = cats
cd.add_series("Base", base)
cd.add_series("Total", tot_s)
cd.add_series("Increase", inc)
ch = add_chart(s, XL_CHART_TYPE.COLUMN_STACKED, cd)
style_chart(ch, False)
clean_axes(ch)
ser = ch.plots[0].series
ser[0].format.fill.background()
ser[0].format.line.fill.background()
for sr, c in ((ser[1], INK), (ser[2], TEAL)):
    sr.format.fill.solid()
    sr.format.fill.fore_color.rgb = rgb(c)
ch.plots[0].gap_width = 40
ch.plots[0].overlap = 100
for sr in (ser[1], ser[2]):
    sr.data_labels.show_value = True
    sr.data_labels.number_format = '#,##0'
    sr.data_labels.number_format_is_linked = False
    sr.data_labels.font.size = Pt(10)
    sr.data_labels.font.bold = True
    sr.data_labels.font.color.rgb = rgb(WHITE)

# ---- treemap (shape built) -------------------------------------------------
def squarify(vals, x, y, w, h):
    """vals sorted desc -> list of (x,y,w,h)"""
    out, vals = [], list(vals)
    tot_v = sum(vals)
    area = w * h
    sc = [v / tot_v * area for v in vals]

    def worst(row, side):
        s_ = sum(row)
        return max(max(side * side * r / (s_ * s_), s_ * s_ / (side * side * r)) for r in row)

    i = 0
    while i < len(sc):
        side = min(w, h)
        row = [sc[i]]
        j = i + 1
        while j < len(sc) and worst(row + [sc[j]], side) <= worst(row, side):
            row.append(sc[j])
            j += 1
        s_ = sum(row)
        if w >= h:
            cw = s_ / h
            yy = y
            for r in row:
                out.append((x, yy, cw, r / cw))
                yy += r / cw
            x += cw
            w -= cw
        else:
            ch_ = s_ / w
            xx = x
            for r in row:
                out.append((xx, y, r / ch_, ch_))
                xx += r / ch_
            y += ch_
            h -= ch_
        i = j
    return out


s = new_slide("Treemap - share of each application at a glance",
              "Global value by end-use application, 2026 - area = market size, label = share",
              "Treemap",
              "Part-to-whole with 7-30 items and long names",
              "Seven applications plus 11 performance functions don't fit a pie; area shows size and colour can show CAGR.",
              "Pie chart with a long legend / long bar list",
              "Colour by CAGR band to show growth in the same view.")
e = ENDUSE.groupby(level=0, sort=False).sum().sort_values("2026", ascending=False)
cg = {n: cagr(e.loc[n, "2026"], e.loc[n, "2036"], 10) for n in e.index}
rects = squarify(e["2026"].tolist(), CH_L, CH_T, CH_W, CH_H)
lo, hi = min(cg.values()), max(cg.values())
for (n, r) in zip(e.index, rects):
    colr, t = ramp(cg[n], lo, hi, (140, 190, 225), (11, 46, 74))
    b = box(s, int(r[0]), int(r[1]), int(r[2]), int(r[3]), colr, WHITE, lw=2)
    box_text(b, [short(n, 30), f"{e.loc[n,'2026']/e['2026'].sum()*100:.1f}%", f"CAGR {cg[n]*100:.1f}%"],
             12 if r[2] > Inches(2) else 9, WHITE if t > 0.35 else INK, True)

# ---- bubble ----------------------------------------------------------------
s = new_slide("Bubble chart - size, growth and scale together",
              "Regions: 2026 value (x) vs 2026-2036 CAGR (y); bubble = 2036 value",
              "Bubble (3 variables)",
              "Prioritising markets: large & fast vs small & fast",
              "Regions differ on both size and growth; one view separates 'big & mature' from 'small & rapidly growing'.",
              "A table of value and CAGR columns",
              "Add a quadrant divider for the 'Opportunity' slide.")
cd = BubbleChartData()
for r in REGIONS.index:
    sr = cd.add_series(short(r, 22))
    sr.add_data_point(round(REGIONS.loc[r, "2026"]),
                      round(cagr(REGIONS.loc[r, "2026"], REGIONS.loc[r, "2036"], 10) * 100, 1),
                      round(REGIONS.loc[r, "2036"]))
ch = add_chart(s, XL_CHART_TYPE.BUBBLE, cd)
style_chart(ch, True, XL_LEGEND_POSITION.BOTTOM)
for i, sr in enumerate(ch.plots[0].series):
    sr.format.fill.solid()
    sr.format.fill.fore_color.rgb = rgb(CAT[i % len(CAT)])
ch.plots[0].bubble_scale = 60
ch.category_axis.has_title = True
ch.category_axis.axis_title.text_frame.text = "2026 value (US$ Mn)"
ch.value_axis.has_title = True
ch.value_axis.axis_title.text_frame.text = "CAGR 2026-2036 (%)"
ch.category_axis.has_major_gridlines = False
ch.value_axis.major_gridlines.format.line.color.rgb = rgb(LIGHT)
ch.category_axis.tick_labels.number_format = '#,##0'
ch.category_axis.tick_labels.number_format_is_linked = False
ch.value_axis.tick_labels.number_format = '0.0"%"'
ch.value_axis.tick_labels.number_format_is_linked = False
labels(ch.plots[0], '#,##0', 9, color=INK)
ch.plots[0].data_labels.show_value = False
ch.plots[0].data_labels.show_series_name = True

# ---- heat-map --------------------------------------------------------------
s = new_slide("Heat-map - end-use × material matrix",
              "Global value 2036 by end-use application and material (US$ Mn)",
              "Colour-scaled matrix",
              "Spotting the hot cells in a two-way table",
              "The dataset already has End-Use × Material: colour shows concentration (Sports × Polyester) at a glance.",
              "A dense numeric table",
              "Same pattern works for Region × Segment CAGR.")
u = USxM.groupby(["Sub-segment", "Sub-segment 1"], sort=False)["2036"].sum().unstack()
u = u.reindex(columns=[m for m in parent.index if m in u.columns])
u = u.loc[u.sum(axis=1).sort_values(ascending=False).index]
nR, nC = u.shape
tbl = s.shapes.add_table(nR + 1, nC + 1, CH_L, CH_T + Inches(0.1), CH_W, Inches(0.55) * (nR + 1)).table
tbl.columns[0].width = Inches(2.35)
for j in range(1, nC + 1):
    tbl.columns[j].width = int((CH_W - Inches(2.35)) / nC)
vmax = float(u.max().max())
hdr = [""] + [short(c, 14) for c in u.columns]
for j, t in enumerate(hdr):
    c = tbl.cell(0, j)
    c.text = t
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(INK)
    r_ = c.text_frame.paragraphs[0].runs
    if r_:
        r_[0].font.size, r_[0].font.bold, r_[0].font.color.rgb = Pt(10), True, rgb(WHITE)
for i, (idx, row) in enumerate(u.iterrows(), start=1):
    c = tbl.cell(i, 0)
    c.text = short(idx, 30)
    c.fill.solid()
    c.fill.fore_color.rgb = rgb(PANEL)
    c.text_frame.paragraphs[0].runs[0].font.size = Pt(10.5)
    c.text_frame.paragraphs[0].runs[0].font.bold = True
    c.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(INK)
    for j, v in enumerate(row, start=1):
        c = tbl.cell(i, j)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        if pd.isna(v) or v == 0:
            c.text = "-"
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(WHITE)
            col = SLATE
        else:
            colr, t = ramp(v, 0, vmax)
            c.text = f"{v:,.0f}"
            c.fill.solid()
            c.fill.fore_color.rgb = rgb(colr)
            col = WHITE if t > 0.45 else INK
        p = c.text_frame.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(11), rgb(col)

# ---- marimekko -------------------------------------------------------------
s = new_slide("Marimekko - size and mix in one chart",
              "Global 2036: column width = end-use size, stack = material share",
              "Marimekko (variable-width 100% stack)",
              "Two dimensions: how big is each segment AND what is inside it",
              "End-use × Material is exactly this: width shows which applications matter, stacks show the fibre mix.",
              "Side-by-side 100% bars that hide segment size",
              "Limit to 5-7 columns and 4-6 stack colours.")
u2 = u.fillna(0)
wtot = u2.sum(axis=1)
X0, Y0, W0, H0 = CH_L + Inches(0.5), CH_T + Inches(0.1), CH_W - Inches(0.6), CH_H - Inches(1.55)
x = X0
mcols = list(u2.columns)
for r_name, rowv in u2.iterrows():
    cw = W0 * (wtot[r_name] / wtot.sum())
    y = Y0
    for k, m in enumerate(mcols):
        share = rowv[m] / wtot[r_name]
        hh = H0 * share
        if hh > 0:
            b = box(s, int(x), int(y), int(cw), int(hh), CAT[k % len(CAT)], WHITE, lw=1)
            if hh > Inches(0.32) and cw > Inches(0.65):
                box_text(b, f"{share*100:.0f}%", 9, WHITE, True)
        y += hh
    tb(s, int(x), int(Y0 + H0 + Inches(0.05)), int(cw), Inches(0.62),
       short(r_name, 20), 8.5, True, INK, PP_ALIGN.CENTER)
    x += cw
for k, m in enumerate(mcols):
    lx = X0 + Inches(1.35) * k
    box(s, int(lx), Y0 + H0 + Inches(0.78), Inches(0.16), Inches(0.16), CAT[k % len(CAT)])
    tb(s, int(lx + Inches(0.18)), Y0 + H0 + Inches(0.72), Inches(1.2), Inches(0.5), short(m, 16), 8.5)

# ---- slope -----------------------------------------------------------------
s = new_slide("Slope chart - who gained, who lost",
              "Regional value (US$ Mn), 2021 vs 2036 - two points per region",
              "Slope chart (2-point line)",
              "Rank change or before/after across 5-10 items",
              "Regions grow at different speeds; East Asia overtakes Western Europe and North America to rank first by 2036.",
              "Two grouped columns per region",
              "Great for 'then vs now' or 'base vs forecast' slides.")
cd = CategoryChartData()
cd.categories = ["2021", "2036"]
for r in REGIONS.index:
    cd.add_series(short(r, 22), [round(REGIONS.loc[r, "2021"]), round(REGIONS.loc[r, "2036"])])
ch = add_chart(s, XL_CHART_TYPE.LINE_MARKERS, cd)
style_chart(ch, True, XL_LEGEND_POSITION.RIGHT)
clean_axes(ch)
for i, sr in enumerate(ch.plots[0].series):
    sr.format.line.color.rgb = rgb(CAT[i % len(CAT)])
    sr.format.line.width = Pt(2.75)
    sr.smooth = False
    sr.marker.format.fill.solid()
    sr.marker.format.fill.fore_color.rgb = rgb(CAT[i % len(CAT)])
labels(ch.plots[0], '#,##0', 9, pos=XL_LABEL_POSITION.RIGHT)

# ---- dumbbell ---------------------------------------------------------------
s = new_slide("Dumbbell - start vs end, ranked",
              "Global value by end-use application: 2026 (hollow) → 2036 (solid), US$ Mn",
              "Dumbbell / connected dot plot",
              "Showing the gap between two periods per category",
              "Length of the bar = absolute growth, so Sports & Activewear vs Footwear can be compared instantly.",
              "Clustered column with two years per group",
              "Add CAGR to the right of the bar.")
e2 = e.sort_values("2036", ascending=True)
vmaxd = float(e2["2036"].max()) * 1.08
LX, TW = CH_L + Inches(2.9), CH_W - Inches(3.6)
n = len(e2)
rowh = (CH_H - Inches(0.6)) / n
for gi in range(6):
    gx = LX + int(TW * gi / 5)
    line(s, gx, CH_T + Inches(0.1), gx, CH_T + CH_H - Inches(0.45), LIGHT, 0.75)
    tb(s, gx - Inches(0.4), CH_T + CH_H - Inches(0.4), Inches(0.8), Inches(0.25),
       f"{vmaxd*gi/5:,.0f}", 9, color=SLATE, align=PP_ALIGN.CENTER)
for i, (nm, rw) in enumerate(e2.iterrows()):
    cy = CH_T + Inches(0.1) + rowh * (i + 0.5)
    x1 = LX + TW * rw["2026"] / vmaxd
    x2 = LX + TW * rw["2036"] / vmaxd
    tb(s, CH_L, int(cy - Inches(0.14)), Inches(2.85), Inches(0.28), short(nm, 32), 11,
       True, INK, PP_ALIGN.RIGHT)
    line(s, int(x1), int(cy), int(x2), int(cy), GREY, 3)
    d = Inches(0.19)
    box(s, int(x1 - d / 2), int(cy - d / 2), d, d, WHITE, BLUE, MSO_SHAPE.OVAL, 2)
    box(s, int(x2 - d / 2), int(cy - d / 2), d, d, BLUE, BLUE, MSO_SHAPE.OVAL, 2)
    tb(s, int(x2 + Inches(0.15)), int(cy - Inches(0.14)), Inches(1.5), Inches(0.28),
       f"{cagr(rw['2026'], rw['2036'], 10)*100:.1f}% CAGR", 10, True, TEAL)

# ---- radar -------------------------------------------------------------------
s = new_slide("Radar - comparing profiles across categories",
              "End-use application mix (% of 2036 value): North America vs East Asia vs Western Europe",
              "Radar / spider",
              "Comparing 2-3 profiles across 5-8 shared axes",
              "Application mix differs by region; the shapes show at once where each region over-indexes.",
              "Three clustered bar charts",
              "Never more than 3 series; 5-8 axes.")
cd = CategoryChartData()
apps = [a for a in e.index]
cd.categories = [short(a, 22) for a in apps]
for rg in ("North America", "East Asia", "Western Europe"):
    sub = D[(D.Region == rg) & (D.Segment == "End-Use Application") & (D.Unit == "Value")]
    sub = sub.groupby("Sub-segment")["2036"].sum()
    cd.add_series(rg, [round(sub.get(a, 0) / sub.sum() * 100, 1) for a in apps])
ch = add_chart(s, XL_CHART_TYPE.RADAR_MARKERS, cd)
style_chart(ch, True, XL_LEGEND_POSITION.BOTTOM)
for sr, c in zip(ch.plots[0].series, (BLUE, CORAL, TEAL)):
    sr.format.line.color.rgb = rgb(c)
    sr.format.line.width = Pt(2.5)
ch.value_axis.tick_labels.number_format = '0"%"'
ch.value_axis.tick_labels.number_format_is_linked = False
ch.value_axis.major_gridlines.format.line.color.rgb = rgb(LIGHT)

# ---- small multiples ---------------------------------------------------------
s = new_slide("Small multiples - one mini-chart per region",
              "Value trend 2021-2036 for every region on the same scale (US$ Mn)",
              "Small multiples / sparkline grid",
              "Many categories that would turn one chart into spaghetti",
              "Seven regions at once with shared axis so the shapes compare fairly.",
              "One line chart with 7+ lines",
              "Use shared axes; highlight one panel for the story.")
gx, gy = 4, 2
cw_, chh = CH_W / gx, CH_H / gy
vmax_all = float(REGIONS[YEARS].max().max()) * 1.05
for i, r in enumerate(REGIONS.index):
    cd = CategoryChartData()
    cd.categories = YEARS
    cd.add_series(r, [round(v) for v in REGIONS.loc[r, YEARS]])
    cx = CH_L + int(cw_ * (i % gx))
    cy = CH_T + int(chh * (i // gx))
    ch = add_chart(s, XL_CHART_TYPE.AREA, cd, cx, cy, int(cw_), int(chh))
    ch.has_legend = False
    ch.has_title = True
    ch.chart_title.text_frame.text = short(r, 20)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
    ch.chart_title.text_frame.paragraphs[0].runs[0].font.bold = True
    ch.font.size = Pt(8)
    ch.value_axis.maximum_scale = vmax_all
    ch.value_axis.minimum_scale = 0
    ch.value_axis.visible = False
    ch.value_axis.has_major_gridlines = False
    ch.category_axis.tick_labels.font.size = Pt(7)
    ch.category_axis.tick_label_position  # keep default
    sr = ch.plots[0].series[0]
    sr.format.fill.solid()
    sr.format.fill.fore_color.rgb = rgb(CAT[i % len(CAT)])
c8 = box(s, CH_L + int(cw_ * 3), CH_T + int(chh), int(cw_), int(chh), PANEL, LIGHT)
box_text(c8, ["Same y-axis on every panel", f"max = US$ {vmax_all:,.0f} Mn"], 11, INK, True)

# ---- tile map ----------------------------------------------------------------
s = new_slide("Tile map - the world in seven tiles",
              "Regions placed roughly where they sit; colour = CAGR 2026-2036, label = 2036 value",
              "Tile-grid map",
              "Geography without drawing a map",
              "Our regions are not countries; tiles keep the geography and let colour carry growth.",
              "Bar chart of regions in arbitrary order",
              "Country-level version: one tile per country (31).")
pos = {"North America": (0, 0), "Western Europe": (2, 0), "Eastern Europe": (3, 0), "East Asia": (4, 0),
       "Latin America": (1, 1), "Middle East & Africa": (3, 1), "South Asia Pacific": (4, 1)}
cg_r = {r: cagr(REGIONS.loc[r, "2026"], REGIONS.loc[r, "2036"], 10) for r in REGIONS.index}
lo, hi = min(cg_r.values()), max(cg_r.values())
tw_, th_ = CH_W / 5, Inches(2.4)
for r, (cx_, cy_) in pos.items():
    colr, t = ramp(cg_r[r], lo, hi, (190, 220, 240), (11, 46, 74))
    b = box(s, CH_L + int(tw_ * cx_) + Inches(0.05), CH_T + Inches(0.3) + int(th_ * cy_) + Inches(0.05),
            int(tw_) - Inches(0.1), int(th_) - Inches(0.1), colr, WHITE)
    box_text(b, [r, f"{cg_r[r]*100:.1f}%", f"US$ {REGIONS.loc[r,'2036']:,.0f} Mn"], 12,
             WHITE if t > 0.4 else INK, True)
tb(s, CH_L, CH_T + Inches(5.2), Inches(7), Inches(0.3), "Darker = faster CAGR", 10, False, SLATE, italic=True)

# ---- KPI cards ---------------------------------------------------------------
s = new_slide("KPI cards with sparkline - headline numbers first",
              "Global performance knitted fabrics at a glance",
              "KPI card + mini column sparkline",
              "Exec summary / first slide of every section",
              "One big number, one delta, one trend - clients read the cards before anything else.",
              "Long text paragraphs of numbers",
              "Max 4 cards per row.")
v26, v36 = GTOT["2026"], GTOT["2036"]
vol = VOL_REG.sum()
cards = [
    ("Market value 2026", f"US$ {v26:,.0f} Mn", f"from US$ {GTOT['2021']:,.0f} Mn in 2021", [GTOT[y] for y in YEARS], BLUE),
    ("Market value 2036", f"US$ {v36:,.0f} Mn", f"{v36/v26:.1f}x vs 2026", [GTOT[y] for y in YEARS], INK),
    ("CAGR 2026-2036", f"{cagr(v26, v36, 10)*100:.1f}%", f"+US$ {v36-v26:,.0f} Mn absolute", [tot[YEARS[i]] / tot[YEARS[i-1]] - 1 for i in range(1, 16)], TEAL),
    ("Volume 2026", f"{vol['2026']:,.0f} Mn Sq. Yd", f"to {vol['2036']:,.0f} Mn Sq. Yd by 2036", [vol[y] for y in YEARS], OCHRE),
]
cw2 = (CH_W - Inches(0.3)) / 2
ch2 = (CH_H - Inches(0.3)) / 2
for i, (t, big, sub, spark, col) in enumerate(cards):
    cx = CH_L + int((cw2 + Inches(0.3)) * (i % 2))
    cy = CH_T + int((ch2 + Inches(0.3)) * (i // 2))
    box(s, cx, cy, int(cw2), int(ch2), PANEL, LIGHT)
    box(s, cx, cy, Inches(0.08), int(ch2), col)
    tb(s, cx + Inches(0.25), cy + Inches(0.15), int(cw2) - Inches(0.4), Inches(0.3), t.upper(), 10, True, SLATE)
    tb(s, cx + Inches(0.25), cy + Inches(0.45), int(cw2) - Inches(0.4), Inches(0.7), big, 26, True, col)
    tb(s, cx + Inches(0.25), cy + Inches(1.15), int(cw2) - Inches(0.4), Inches(0.3), sub, 11, False, INK)
    mn, mx = min(spark), max(spark)
    bw = (cw2 - Inches(0.6)) / len(spark)
    for k, v in enumerate(spark):
        hh = Inches(0.55) * (0.15 + 0.85 * (v - mn) / (mx - mn if mx != mn else 1))
        box(s, cx + Inches(0.25) + int(bw * k), cy + int(ch2) - Inches(0.2) - int(hh),
            int(bw) - Inches(0.03), int(hh), col if k >= 5 else GREY)

prs.save(OUT)
print("saved", OUT, "slides:", len(prs.slides))
