"""Put the Volume panels back on a country's segment pages.

    python restore_country_volume_pages.py "deck.pptx" "source.pptx" "Germany" ["out.pptx"]

The house-style step (deck_style.country_volume_off) drops the Volume table /
chart from a country's "Segmental Analysis", "Tables" and "YoY Analysis"
pages.  ``source`` is a country-only build of the same workbook and settings
made with FMI_DECK_STYLE=off, so those pages still carry Volume.  For every
page of ``country`` whose title exists in both decks and whose content
differs, the slide's shapes are replaced by the source slide's shapes
(charts, tables, pictures, padlocks included).  Slide order and count do not
change, and the other slides are not touched.  A locked deck stays locked:
the copied tables / charts are the locked ones.
"""
import copy
import io
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation                              # noqa: E402
from pptx.opc.constants import RELATIONSHIP_TYPE as RT     # noqa: E402
from pptx.parts.chart import ChartPart                     # noqa: E402
from fmi_report import deck_style as D                     # noqa: E402

R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
C = "{http://schemas.openxmlformats.org/drawingml/2006/chart}"
PAGE = re.compile(r"(Segmental Analysis|Market Tables|Tables|YoY Analysis)\s+\S")


def _copy_chart(src_sl, dst_sl, rid):
    sp = src_sl.part.related_part(rid)
    pkg = dst_sl.part.package
    new = ChartPart.load(pkg.next_partname("/ppt/charts/chart%d.xml"),
                         sp.content_type, pkg, sp.blob)
    try:
        xl = sp.chart_workbook.xlsx_part
        if xl is not None:
            new.chart_workbook.update_from_xlsx_blob(xl.blob)
    except Exception:
        pass
    return dst_sl.part.relate_to(new, RT.CHART)


def _copy_shape(src_sl, dst_sl, sh):
    el = copy.deepcopy(sh._element)
    for node in list(el.iter()):
        rid = node.get(R + "embed")
        if rid:
            blob = src_sl.part.related_part(rid).blob
            _, new = dst_sl.part.get_or_add_image_part(io.BytesIO(blob))
            node.set(R + "embed", new)
        if node.tag == C + "chart" and node.get(R + "id"):
            node.set(R + "id", _copy_chart(src_sl, dst_sl, node.get(R + "id")))
        if node.tag.endswith("hlinkClick"):
            node.getparent().remove(node)
    dst_sl.shapes._spTree.append(el)


def _fingerprint(sl):
    return sorted((s.name, s.left, s.top, s.width, s.height) for s in sl.shapes)


def main(deck, source, country, out):
    prs, src = Presentation(str(deck)), Presentation(str(source))
    by_title = {}
    for sl in src.slides:
        by_title.setdefault(D.title_of(sl), sl)
    done, same, missing = [], 0, []
    for i, sl in enumerate(prs.slides, 1):
        t = D.title_of(sl)
        if not t.startswith(country) or not PAGE.search(t):
            continue
        s = by_title.get(t)
        if s is None:
            missing.append((i, t))
            continue
        if _fingerprint(sl) == _fingerprint(s):
            same += 1
            continue
        tree = sl.shapes._spTree
        for el in [e for e in tree if e.tag.endswith(("}sp", "}pic", "}graphicFrame",
                                                      "}grpSp", "}cxnSp"))]:
            tree.remove(el)
        for sh in s.shapes:
            _copy_shape(s, sl, sh)
        done.append((i, t))
    prs.save(str(out))
    print(f"{Path(deck).name}: {len(done)} pages restored, {same} already complete, "
          f"{len(missing)} without a source page")
    for i, t in done:
        print("  restored", i, t[:90].encode("ascii", "replace").decode())
    for i, t in missing:
        print("  no source page:", i, t[:90].encode("ascii", "replace").decode())


if __name__ == "__main__":
    a = sys.argv
    main(a[1], a[2], a[3], a[4] if len(a) > 4 else a[1])
