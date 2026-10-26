"""Take every Volume chart, graph and table out of a finished deck (output only).

    python strip_volume_content.py "deck.pptx" ["out.pptx"]

For a value-only market.  Works on the .pptx, the engine is not involved:
  * a panel whose heading names Volume ("Global Volume By Region (Tons)",
    "U.S. Volume YoY, By Product Type (%)") is removed with everything inside
    its card; the panel left on the page moves up / takes the full width, as
    the house style does for the country pages;
  * a page left with a heading but no body is reported, never deleted;
  * the words 'volume and value' / 'Volume & Value' in headings, sub-headings
    and the agenda become 'value' / 'Value'; an agenda bullet that is only
    about volume is dropped.
"""
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from pptx import Presentation                       # noqa: E402
from fmi_report import deck_style as D              # noqa: E402

PANEL = re.compile(r"\bVolume\b.*(\(Tons\)|\(Units\)|\(MT\)|\(%\)|\bYoY\b|"
                   r"By |Share|Shares)|\(Tons\)|\(Units\)")


def volume_panel_titles(slide):
    out = []
    title = D.title_of(slide)
    for s in slide.shapes:
        if D._shape_type(s) != 17 or not s.has_text_frame:
            continue
        t = s.text_frame.text.strip()
        if not t or t == title or s.top / D.PX < 100:
            continue
        if re.search(r"\bVolume\b", t) and PANEL.search(t) \
                and not re.search(r"\bValue\b", t.split("(")[0]):
            out.append(s)
    return out


def strip_panels(slide):
    n = 0
    while True:
        ts = volume_panel_titles(slide)
        if not ts:
            return n
        cb = D.drop_panel(slide, ts[0])
        if cb is None:
            D.remove_shape(ts[0])
            continue
        n += 1
        others = [s for s in slide.shapes if s.shape_type == 1
                  and s.name.startswith("Rounded Rectangle")
                  and s.width / D.PX > 300 and s.top / D.PX > cb[3] - 5]
        if others:
            top2 = min(s.top / D.PX for s in others)
            D.shift_up(slide, top2, top2 - cb[1])
        else:
            D.widen_row(slide, cb[1], 1250)


def tidy_words(slide):
    for s in slide.shapes:
        if not s.has_text_frame or D._shape_type(s) != 17:
            continue
        tf = s.text_frame
        # an agenda / divider bullet that is only about volume
        for p in list(tf.paragraphs):
            tx = "".join(r.text for r in p.runs)
            if tx.lstrip().startswith("➤") and re.search(r"\bVolume\b", tx) \
                    and not re.search(r"\bValue\b", tx):
                p._p.getparent().remove(p._p)
        for p in tf.paragraphs:
            for r in p.runs:
                new = re.sub(r"Volume (and|&) Value", "Value", r.text)
                new = re.sub(r"(?i)\bvolume (and|&) value\b", "value", new)
                if new != r.text:
                    r.text = new


def fill_body(slide, bottom=688, cap=2.2):
    """A page whose volume panel went has one card at the top and a blank
    lower half: stretch that card (and what is in it) down the page."""
    cards = [s for s in slide.shapes if s.shape_type == 1
             and s.name.startswith("Rounded Rectangle")
             and s.width / D.PX > 300]
    if len(cards) != 1:
        return False
    card = cards[0]
    cb = D.box(card)
    top0, h = cb[1], cb[3] - cb[1]
    if cb[3] > 470 or h <= 0:
        return False
    f = min((bottom - top0) / h, cap)
    if f < 1.15:
        return False
    for sh in list(slide.shapes):
        if sh is not card and not D.inside(D.box(sh), cb):
            continue
        t, ht = sh.top / D.PX, sh.height / D.PX
        if sh.has_text_frame and D._shape_type(sh) == 17:   # text: move only
            sh.top = D.px(top0 + (t - top0) * f)
        elif D._shape_type(sh) == 13:                       # picture: re-centre
            sh.top = D.px(top0 + (t + ht / 2 - top0) * f - ht / 2)
        else:                                               # card, chart, table
            sh.top = D.px(top0 + (t - top0) * f)
            if getattr(sh, "has_table", False) and sh.has_table:
                for r in sh.table.rows:
                    r.height = int(r.height * f)
            else:
                sh.height = int(sh.height * f)
    return True


def main(deck, out):
    prs = Presentation(str(deck))
    dropped = 0
    filled = 0
    for sl in prs.slides:
        n = strip_panels(sl)
        dropped += n
        tidy_words(sl)
        if n or re.search(r"Trends, By|YoY Analysis", D.title_of(sl)):
            filled += bool(fill_body(sl))
    prs.save(str(out))
    empty = [(i, D.title_of(s)) for i, s in enumerate(prs.slides, 1)
             if D.title_of(s) and D.empty_body(s)]
    print(f"{dropped} volume panel(s) removed; {filled} page(s) stretched; {len(empty)} titled page(s) "
          "left without a body")
    for i, t in empty:
        print("  empty:", i, t[:80].encode("ascii", "replace").decode())


if __name__ == "__main__":
    a = sys.argv
    main(a[1], a[2] if len(a) > 2 else a[1])
