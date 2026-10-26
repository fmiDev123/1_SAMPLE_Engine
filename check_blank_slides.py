"""Flag slides with no body: no chart, table or picture and almost no text.

    python check_blank_slides.py "deck.pptx"
"""
import sys
from pptx import Presentation


def _type(sh):
    try:
        return sh.shape_type
    except NotImplementedError:          # a shape python-pptx cannot classify
        return None


def walk(shs):
    for sh in shs:
        if _type(sh) == 6:
            yield from walk(sh.shapes)
        else:
            yield sh


p = Presentation(sys.argv[1])
out = []
for i, s in enumerate(p.slides, 1):
    vis, t = 0, []
    for sh in walk(s.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip():
            t.append(sh.text_frame.text.strip())
        if getattr(sh, "has_table", False) and sh.has_table:
            vis += 1
            t += [c.text for r in sh.table.rows for c in r.cells]
        if getattr(sh, "has_chart", False) and sh.has_chart:
            vis += 1
        if _type(sh) == 13:
            vis += 1
    body = [x for x in t if "All Rights Reserved" not in x and "Source@" not in x]
    chars = sum(len(x) for x in body)
    if vis == 0 and chars < 200:
        out.append((i, chars, (body[0] if body else "")[:60].replace("\n", " ")))
for o in out:
    print(o)
print(len(out), "sparse slides of", len(p.slides))
sys.exit(1 if out else 0)
