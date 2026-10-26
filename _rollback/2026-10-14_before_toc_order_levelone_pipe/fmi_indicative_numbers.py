#!/usr/bin/env python3
"""
fmi_indicative_numbers.py
=========================
Fill a Formatted Dataset skeleton (every cell 'XX') with INDICATIVE numbers, so
a sample deck can be built for a market the research team has not sized yet.

    python fmi_indicative_numbers.py REPORT_OR_RAW.xlsx -o OUT.xlsx
            [--market "Packaged Muffins"] [--global-value 10900] [--price 4100]

THE FIGURES ARE NOT RESEARCH.  They are generated, deterministic (the same
input always gives the same numbers) and internally consistent:

  * every segment of a geography sums to the same total, Value and Volume;
  * a region is the sum of its countries, the top geography the sum of its
    regions;
  * cross segments (A x B) are fitted to the A and B shares, so the matrix and
    the single-segment pages agree;
  * Volume is Value / price, with a price that differs by geography and by
    sub-segment, so the derived prices and indices are not flat.

Anything the report DOES carry as a real number is kept as it is - a leaf row
with numbers in every year is never overwritten, and when the base-year
totals are given (--global-value, --price) the model is pinned to them.

Built on fmi_formatted_dataset (read / write / QC).
"""
from __future__ import annotations

import argparse
import hashlib
import re
import sys
from collections import OrderedDict, defaultdict

import fmi_formatted_dataset as F

BASE_YEAR = 2026

# rough regional weights of a global consumer-goods market; any other region
# name gets a stable pseudo-random weight
_REGION_W = {"north america": 0.33, "europe": 0.27, "east asia": 0.13,
             "south asia & pacific": 0.12, "latin america": 0.08,
             "middle east & africa": 0.07}


# Average-price tier of a region against the world average (= 1.00), used
# whenever the research team has not given prices: North America highest,
# Europe above the average, Latin America and MEA slightly below, Asia Pacific
# last (China and India pull it down).  A region not listed gets 0.95.
_REGION_PRICE = {
    "north america": 1.35,
    "europe": 1.15, "western europe": 1.20, "eastern europe": 0.98,
    "latin america": 0.92,
    "middle east & africa": 0.88, "middle east and africa": 0.88,
    "mea": 0.88, "middle east": 0.90, "africa": 0.82,
    "east asia": 0.74, "south asia & pacific": 0.66,
    "asia pacific": 0.70, "apac": 0.70, "asia": 0.70, "oceania": 0.95,
}
# Annual price drift of a region (the world's blended price rises ~0.6% a year): North
# America and Europe climb faster than the world, Latin America and MEA sit
# close to it, East Asia and South Asia & Pacific lag it - so a "price
# variation against Global" chart shows some regions up and some down.
_REGION_PRICE_DRIFT = {
    "north america": 0.0140, "europe": 0.0120,
    "western europe": 0.0125, "eastern europe": 0.0105,
    "latin america": 0.0060,
    "middle east & africa": 0.0045, "middle east and africa": 0.0045,
    "mea": 0.0045,
    "east asia": 0.0010, "south asia & pacific": -0.0010,
    "asia pacific": 0.0000, "apac": 0.0000,
}
_DEFAULT_PRICE_DRIFT = 0.0090
# the lowest-price markets inside any region
_LOW_PRICE_COUNTRIES = {"india", "china", "indonesia", "vietnam", "philippines",
                        "pakistan", "bangladesh", "thailand", "malaysia"}


def _h(*parts) -> float:
    """Stable pseudo-random number in [0, 1) from the parts."""
    d = hashlib.md5("|".join(map(str, parts)).encode("utf-8")).hexdigest()
    return int(d[:8], 16) / 0xFFFFFFFF


def _is_num(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _round_alloc(total, weights, dec=3):
    """Split ``total`` over ``weights`` rounded to ``dec`` places so the parts
    add up EXACTLY to round(total, dec) - the largest part takes the slack."""
    s = sum(weights) or 1.0
    parts = [round(total * w / s, dec) for w in weights]
    gap = round(round(total, dec) - sum(parts), dec)
    if parts and gap:
        k = max(range(len(parts)), key=lambda i: parts[i])
        parts[k] = round(parts[k] + gap, dec)
    return parts


END_YEAR = 2036


def headline_figures(value_2025=None, value_2026=None, value_2036=None,
                     cagr=None, base_year=BASE_YEAR, end_year=END_YEAR):
    """Complete the Global headline set from whatever was given.

    Any of: the base-year (2026) value, the end-year (2036) value, the CAGR
    over base -> end (a fraction, 0.044 = 4.4%) and the 2025 value.  The
    rest follows from the basic formulas

        CAGR      = (V_end / V_base) ** (1 / n) - 1          n = end - base
        V_end     = V_base * (1 + CAGR) ** n
        V_base    = V_end / (1 + CAGR) ** n
        V_2025    = V_base / (1 + CAGR)                      (when not given)
        absolute $ opportunity = V_end - V_base              (2026 -> 2036)
                               = V_end - V_2025              (2025 -> 2036,
                                                              the classic text)

    Returns a dict (None where it cannot be worked out) - never overwrites a
    figure that was given.
    """
    n = end_year - base_year
    v25, v26, v36, g = value_2025, value_2026, value_2036, cagr
    if g is not None and g > 1:          # 4.4 given as a percentage
        g = g / 100.0
    if v26 is None and v36 is not None and g is not None:
        v26 = v36 / (1 + g) ** n
    if v26 is None and v25 is not None and g is not None:
        v26 = v25 * (1 + g)
    if v36 is None and v26 is not None and g is not None:
        v36 = v26 * (1 + g) ** n
    if g is None and v26 and v36:
        g = (v36 / v26) ** (1.0 / n) - 1
    if v25 is None and v26 is not None and g is not None:
        v25 = v26 / (1 + g)
    out = {"value_2025": v25, "value_2026": v26, "value_2036": v36,
           "cagr": g, "abs_opportunity": None, "abs_opportunity_2025": None}
    if v26 is not None and v36 is not None:
        out["abs_opportunity"] = v36 - v26
    if v25 is not None and v36 is not None:
        out["abs_opportunity_2025"] = v36 - v25
    return out


def read_snapshot(path):
    """The Global headline numbers an Excel Report already carries on its
    'Executive Summary' sheet ('Global Market Snapshot': Market Size, Value
    row with 2021 / 2026 / 2031 / 2036 and the CAGR), or {} when they are
    still placeholders."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = next((w for w in wb.worksheets
                   if w.title.strip().lower().startswith("executive")), None)
        if ws is None:
            return {}
        head = None
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() if c is not None else "" for c in row]
            if cells and cells[0].lower() == "indicator":
                head = cells
                continue
            if head and cells and cells[0].lower().startswith(
                    "market size, value"):
                got = {}
                for h, v in zip(head[1:], row[1:]):
                    try:
                        f = float(str(v).replace(",", "").replace("%", ""))
                    except (TypeError, ValueError):
                        continue
                    if h.lower().startswith("cagr"):
                        got["cagr"] = f
                    elif h.strip().isdigit():
                        got["value_" + h.strip()] = f
                return got
    except Exception:
        pass
    return {}


def volume_factor(unit) -> float:
    """value (US$ Mn) * factor / price = volume, for the volume unit given:
    tonnes / kg / litres -> price per tonne etc. (1e6); 'Mn Units' -> price
    per unit (1); 'Thousand Units' -> 1e3."""
    u = str(unit or "").lower()
    if re.search(r"\b(mn|million|m)\b", u) and re.search(
            r"unit|piece|pack|kit|pair|bottle|device|set", u):
        return 1.0
    if re.search(r"thousand|000|\bk\b", u) and re.search(
            r"unit|piece|pack|kit|pair|bottle|device|set", u):
        return 1e3
    return 1e6


class Model:
    def __init__(self, ds, global_value, price, dec=3, cagr=None):
        self.ds = ds
        self.years = ds.years
        self.n = len(self.years)
        self.dec = dec
        self.global_value = float(global_value)
        self.cagr = cagr          # pins the top geography to base * (1+g)^d
        self.vol_factor = volume_factor(ds.units.get(F.VOLUME))
        self.price0 = float(price)
        self.top = ds.top_geo()
        self.children = ds.geo_children
        self.leaf_geos = [g for g in ds.geo_order if g not in self.children]
        self.T = {}          # (geo, metric) -> [n] totals
        self._weights()

    # ---- geography totals ------------------------------------------------
    def _weights(self):
        ds = self.ds
        regions = list(self.children.get(self.top, []))
        rw = {}
        for r in regions:
            rw[r] = _REGION_W.get(r.casefold(), 0.04 + 0.12 * _h("rw", r))
        tot = sum(rw.values()) or 1.0
        self.geo_w = {}
        for r in regions:
            kids = self.children.get(r, [])
            if not kids:
                self.geo_w[r] = rw[r] / tot
                continue
            cw = {}
            for k, c in enumerate(kids):
                w = 1.0 / (k + 1) ** 1.15
                if c.casefold().startswith(("rest of", "other")):
                    w *= 0.55
                cw[c] = w
            ct = sum(cw.values()) or 1.0
            for c in kids:
                self.geo_w[c] = rw[r] / tot * cw[c] / ct

    def _price_index(self):
        """Relative price of every leaf geography: its region's tier, a small
        deterministic country variation, low-price countries pulled down -
        scaled so the top geography's average price is exactly ``price``."""
        parent = {c: r for r, cs in self.children.items() for c in cs}
        raw = {}
        for g in self.leaf_geos:
            reg = parent.get(g, g)
            if reg == self.top:
                reg = g
            base = _REGION_PRICE.get(reg.casefold(), 0.95)
            if g != reg:
                base *= 0.94 + 0.12 * _h("pi", g)
                if g.casefold() in _LOW_PRICE_COUNTRIES:
                    base *= 0.88
            raw[g] = base
        w = {g: self.geo_w[g] for g in raw if g in self.geo_w}
        sw = sum(w.values())
        den = sum(w[g] / raw[g] for g in w)
        k = (sw / den) if den else 1.0       # harmonic, value-weighted mean
        self.price_idx = {g: raw[g] / k for g in raw}
        # per-geography annual price drift: the region's, plus a small
        # deterministic country spread
        self.price_drift = {}
        for g in self.leaf_geos:
            reg = parent.get(g, g)
            if reg == self.top:
                reg = g
            dr = _REGION_PRICE_DRIFT.get(reg.casefold(), _DEFAULT_PRICE_DRIFT)
            if g != reg:
                dr += 0.0016 * (_h("pd", g) - 0.5) * 2
            self.price_drift[g] = dr

    def totals(self):
        yrs = self.years
        self._price_index()
        for g in self.leaf_geos:
            w = self.geo_w.get(g)
            if w is None:
                continue
            v0 = self.global_value * w
            gf = 0.052 + 0.016 * (_h("g", g) - 0.5) * 2
            gh = gf - 0.008
            val, vol = [], []
            pidx = self.price_idx[g]
            for y in yrs:
                d = y - BASE_YEAR
                v = v0 * ((1 + gf) ** d if d >= 0 else 1.0 / (1 + gh) ** (-d))
                p = self.price0 * pidx * ((1 + self.price_drift[g]) ** d)
                val.append(round(v, self.dec))
                vol.append(round(v * self.vol_factor / p, self.dec))
            self.T[(g, F.VALUE)], self.T[(g, F.VOLUME)] = val, vol
        if self.cagr is not None:
            # pin the world to the headline path: every leaf is scaled by the
            # same per-year factor, so regions and the top still add up
            for k, y in enumerate(yrs):
                tgt = self.global_value * (1 + self.cagr) ** (y - BASE_YEAR)
                tot = sum(self.T[(g, F.VALUE)][k] for g in self.leaf_geos
                          if (g, F.VALUE) in self.T)
                f = tgt / tot if tot else 1.0
                for g in self.leaf_geos:
                    if (g, F.VALUE) in self.T:
                        for m in (F.VALUE, F.VOLUME):
                            self.T[(g, m)][k] = round(self.T[(g, m)][k] * f,
                                                      self.dec)
        # aggregates bottom-up: regions, then the top
        for agg in [g for g in self.children if g != self.top] + [self.top]:
            kids = self.children.get(agg, [])
            for m in F.METRICS:
                parts = [self.T[(k, m)] for k in kids if (k, m) in self.T]
                if parts:
                    self.T[(agg, m)] = [round(sum(c), self.dec)
                                        for c in zip(*parts)]

    # ---- shares within a segment ------------------------------------------
    def _raw_w(self, seg, leaf, geo, y, metric):
        base = (0.35 + _h("b", seg, leaf)) ** 1.7
        tilt = 1.0 + 0.5 * (_h("t", geo, seg, leaf) - 0.5)
        drift = 1.0 + 0.06 * (_h("d", seg, leaf) - 0.5) * (y - BASE_YEAR)
        w = base * tilt * max(0.2, drift)
        if metric == F.VOLUME:           # cheaper sub-segments weigh more in MT
            w /= 0.7 + 0.7 * _h("p", seg, leaf)
        return w

    def fill(self):
        ds, yrs, n = self.ds, self.years, self.n
        by_geo = defaultdict(lambda: defaultdict(list))
        for key in ds.rows:
            g, seg, a, b, c, m = key
            by_geo[(g, m)][seg].append(key)
        out = {}
        shares = {}                      # (geo, metric, seg) -> {name: [n]}
        order = ([g for g in self.leaf_geos]
                 + [g for g in self.children if g != self.top] + [self.top])

        def level_shares(g, m, seg, keys, tot):
            lv = {}
            for k in keys:
                for nm in {k[2], k[3], k[4]}:
                    cur = lv.setdefault(nm, [0.0] * n)
                    for i in range(n):
                        cur[i] += out[k][i] / (tot[i] or 1.0)
            shares[(g, m, seg)] = lv

        for g in order:
            kids = self.children.get(g, [])
            for m in F.METRICS:
                tot = self.T.get((g, m))
                if tot is None:
                    continue
                segs = by_geo.get((g, m), {})
                # 1. the geography breakdown rows: each child's own total
                for seg, keys in segs.items():
                    if (g, seg) in ds.geo_segments:
                        for key in keys:
                            if (key[2], m) in self.T:
                                out[key] = list(self.T[(key[2], m)])
                # 2. an aggregate geography is the SUM of its children, row by
                #    row; only rows no child carries are allocated (below)
                if kids:
                    for seg, keys in segs.items():
                        if (g, seg) in ds.geo_segments:
                            continue
                        for k in keys:
                            ck = [(c,) + k[1:] for c in kids]
                            have = [out[x] for x in ck if x in out]
                            if have:
                                out[k] = [round(sum(c), self.dec)
                                          for c in zip(*have)]
                # 3. plain segments (leaf geographies in full, aggregates for
                #    the rows still open)
                for seg, keys in segs.items():
                    if (g, seg) in ds.geo_segments or seg in ds.cross_segments:
                        continue
                    todo = [k for k in keys if k not in out]
                    if todo:
                        done = [k for k in keys if k in out]
                        rem = [tot[i] - sum(out[k][i] for k in done)
                               for i in range(n)]
                        ws = {k: [self._raw_w(seg, k[4], g, y, m)
                                  for y in yrs] for k in todo}
                        cols = [_round_alloc(max(rem[i], 0.0),
                                             [ws[k][i] for k in todo],
                                             self.dec) for i in range(n)]
                        for j, k in enumerate(todo):
                            out[k] = [cols[i][j] for i in range(n)]
                    level_shares(g, m, seg, keys, tot)
                # 4. cross segments, fitted to the A and B shares
                for seg, keys in segs.items():
                    if seg not in ds.cross_segments:
                        continue
                    todo = [k for k in keys if k not in out]
                    if not todo:
                        continue
                    done = [k for k in keys if k in out]
                    rem = [max(tot[i] - sum(out[k][i] for k in done), 0.0)
                           for i in range(n)]
                    dims = [d.strip() for d in seg.split(F.CROSS_SEP.strip())]
                    sa = shares.get((g, m, dims[0]), {})
                    sb = (shares.get((g, m, dims[1]), {})
                          if len(dims) > 1 else {})
                    a_names = list(OrderedDict.fromkeys(k[3] for k in todo))
                    b_names = list(OrderedDict.fromkeys(k[4] for k in todo))
                    cols = []
                    for i in range(n):
                        ma = {a: sa.get(a, [1.0 / len(a_names)] * n)[i]
                              for a in a_names}
                        mb = {b: sb.get(b, [1.0 / len(b_names)] * n)[i]
                              for b in b_names}
                        za, zb = sum(ma.values()) or 1.0, sum(mb.values()) or 1.0
                        ma = {a: v / za for a, v in ma.items()}
                        mb = {b: v / zb for b, v in mb.items()}
                        cell = {(k[3], k[4]): ma[k[3]] * mb[k[4]]
                                * (0.75 + 0.5 * _h("x", g, seg, k[3], k[4]))
                                for k in todo}
                        for _ in range(8):            # fit to both marginals
                            for a in a_names:
                                t = sum(v for (aa, _b), v in cell.items()
                                        if aa == a) or 1.0
                                for kk in cell:
                                    if kk[0] == a:
                                        cell[kk] *= ma[a] / t
                            for b in b_names:
                                t = sum(v for (_a, bb), v in cell.items()
                                        if bb == b) or 1.0
                                for kk in cell:
                                    if kk[1] == b:
                                        cell[kk] *= mb[b] / t
                        cols.append(_round_alloc(
                            rem[i], [cell[(k[3], k[4])] for k in todo],
                            self.dec))
                    for j, k in enumerate(todo):
                        out[k] = [cols[i][j] for i in range(n)]
        return out


def fill_indicative(ds, global_value=10900.0, price=4100.0, cagr=None):
    """Replace placeholder cells of ``ds`` with indicative numbers (rows that
    already carry real numbers in every year are kept as given)."""
    m = Model(ds, global_value, price, cagr=cagr)
    m.totals()
    new = m.fill()
    kept = 0
    for key, vals in ds.rows.items():
        if all(_is_num(v) for v in vals) and any(
                float(v) != 0 for v in vals):
            kept += 1                    # a real row: never overwritten
            continue                     # (an all-zero row is a placeholder)
        if key in new:
            ds.rows[key] = new[key]
    ds.notes.append(f"indicative figures generated for placeholder rows "
                    f"({kept} real rows kept)")
    return ds


def write_companies(names, out_xlsx):
    """A 'Company Mapping' sheet the deck's company pages read."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Company Mapping"
    ws.append(["Company Name"])
    for n in names:
        ws.append([n])
    wb.save(out_xlsx)
    return out_xlsx


def build_from_segmentation(seg_path, out_dir=None, market=None, price=None):
    """Segmentation brief -> indicative dataset (+ marker + companies sheet).

    Returns the dataset path.  The headline figures in the brief (2026 / 2036
    value, CAGR) pin the model; whatever is missing is worked out
    (headline_figures).  No regions in the brief -> the standard regional
    taxonomy."""
    from pathlib import Path
    seg = Path(seg_path)
    out_dir = Path(out_dir) if out_dir else seg.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    ds = F.read_segmentation(seg, market=market)
    h = headline_figures(**{k: v for k, v in (ds.headline or {}).items()
                            if k != "volume_unit"})
    pinned = h["value_2026"] is not None and h["cagr"] is not None
    p = price
    if p is None:
        p = 4100.0 if volume_factor(ds.units.get(F.VOLUME)) > 1 else 10.0
    fill_indicative(ds, h["value_2026"] if h["value_2026"] is not None
                    else 10900.0, p, cagr=h["cagr"] if pinned else None)
    stem = re.sub(r"[^A-Za-z0-9]+", "_", f"Global {ds.market} Market").strip("_")
    out = out_dir / f"{stem}_Formatted_Dataset_ind.xlsx"
    F.write_formatted_dataset(ds, out)
    with open(str(out) + ".indicative", "w", encoding="utf-8") as fh:
        fh.write("Figures in this dataset are generated examples, not "
                 "research. Decks built from it withhold every figure.\n")
    if ds.companies:
        write_companies(ds.companies,
                        out_dir / f"{stem}_Companies.xlsx")
    return out


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("input")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--market")
    ap.add_argument("--global-value", type=float, default=None,
                    help="top-geography base-year (2026) value, US$ Mn "
                         "(default: the report's Executive Summary snapshot "
                         "if it carries numbers, else 10900)")
    ap.add_argument("--value-2025", type=float, default=None,
                    help="2025 value; worked out as 2026 / (1 + CAGR) "
                         "when not given")
    ap.add_argument("--value-2036", type=float, default=None,
                    help="2036 value; with --global-value gives the CAGR")
    ap.add_argument("--cagr", type=float, default=None,
                    help="CAGR 2026-2036 in percent (4.4) or as a fraction")
    ap.add_argument("--price", type=float, default=None,
                    help="top-geography base-year average price, US$/MT "
                         "(US$ per unit when the volume unit is 'Mn Units'; "
                         "default 4100 / 10)")
    a = ap.parse_args(argv)
    ds = F.load_any(a.input, market=a.market)
    snap = read_snapshot(a.input) if str(a.input).lower().endswith(
        (".xlsx", ".xlsm")) else {}
    for k, v in (getattr(ds, "headline", None) or {}).items():
        if k != "volume_unit" and v is not None and snap.get(k) is None:
            snap[k] = v          # the segmentation's own parameters
    given = dict(value_2025=a.value_2025, value_2026=a.global_value,
                 value_2036=a.value_2036, cagr=a.cagr)
    for key in list(given):
        if given[key] is None and snap.get(key) is not None:
            given[key] = snap[key]
    h = headline_figures(**given)
    pinned = h["value_2026"] is not None and h["cagr"] is not None
    if pinned:
        print("Headline figures (given or worked out):")
        print(f"  2025 value            {h['value_2025']:,.1f}")
        print(f"  2026 value            {h['value_2026']:,.1f}")
        print(f"  2036 value            {h['value_2036']:,.1f}")
        print(f"  CAGR 2026-2036        {h['cagr'] * 100:.2f}%")
        print(f"  Absolute $ opp. 2026->2036  {h['abs_opportunity']:,.1f}")
        print(f"  Absolute $ opp. 2025->2036  "
              f"{h['abs_opportunity_2025']:,.1f}")
    elif h["value_2026"] is not None:
        a.global_value = h["value_2026"]
    if a.price is None:
        a.price = 4100.0 if volume_factor(ds.units.get(F.VOLUME)) > 1 else 10.0
    fill_indicative(ds, h["value_2026"] if h["value_2026"] is not None
                    else 10900.0, a.price,
                    cagr=h["cagr"] if pinned else None)
    issues = F.run_qc(ds)
    F.write_formatted_dataset(ds, a.out)
    # a marker beside the workbook: decks built from it lock EVERY figure
    with open(str(a.out) + ".indicative", "w", encoding="utf-8") as fh:
        fh.write("Figures in this dataset are generated examples, not "
                 "research. Decks built from it withhold every figure.\n")
    errs = [i for i in issues if "ERROR" in str(i)]
    print(f"Wrote {a.out}  |  QC: {len(errs)} error(s), "
          f"{len(issues) - len(errs)} other message(s)")
    for i in issues[:15]:
        print("  ", i)
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main())
