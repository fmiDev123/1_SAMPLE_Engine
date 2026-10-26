#!/usr/bin/env python3
"""
fmi_formatted_dataset.py
========================
Build the FMI **Formatted Dataset** workbook

    Home | Master Sheet | Value | Volume

from whatever the project has on hand:

  1. an FMI **Excel Report** workbook          ('Market Data' sheet + optional 'Cross-sectional')
  2. an FMI **Raw Data** workbook              (Region | Segment | Sub-segment | Sub-segment 1 |
                                                Sub-segment 2 | Unit | 2021 ... 2036 [| CAGR/YoY])
     (an existing Formatted Dataset's 'Master Sheet' is read the same way)
  3. **in-memory records**                     (list of dicts, or pandas DataFrame.to_dict('records'))
  4. a **segmentation only** (taxonomy + geographies) -> 'XX' skeleton for sample decks

Output layout (matches the FMI formatted dataset, e.g. Europe Pet Cosmetic Products):

  Master Sheet  title merged A1:<last>4, header on row 6:
                Region | Segment | Sub-segment | Sub-segment 1 | Sub-segment 2 | Unit | years...
                Unit column holds the METRIC NAME ('Value' / 'Volume'), not 'US$ Mn' / 'MT';
                the units live in the sheet title. (--unit-style unit switches back to US$ Mn / MT.)
                Leaf rows only. Shallow leaves repeat their name across the sub-segment columns.
                Geography breakdown rows: Segment 'Region' (under the top geography) and
                'Country' (under each region), name repeated in all three sub-segment columns.
                Order: geography (source order) -> Value block then Volume block ->
                geography segment, segments (source order), cross segments.
  Value/Volume  pivot-style outline pasted as values: 'Unit' filter on row 19, 'Row Labels'
                header on row 21, data from row 22. Levels: geography > segment > sub-segment >
                sub-segment 1 > [sub-segment 2 only where it differs from sub-segment 1].
                Sorted A-Z like a pivot ('Other...' / 'Rest of...' last); aggregate geographies
                first in source order, then countries A-Z. Grand Total = top geography.
  Home          navigation buttons (hyperlinked shapes) to the three data sheets.

CLI
---
    python fmi_formatted_dataset.py INPUT.xlsx [-o OUTPUT.xlsx] [--market "Packaged Muffins"]
                                    [--unit-style metric|unit] [--value-decimals 3]
                                    [--volume-decimals 3] [--strict]

Library
-------
    from fmi_formatted_dataset import load_any, write_formatted_dataset, run_qc
    ds = load_any("Report.xlsx")                    # auto-detects report vs raw data
    issues = run_qc(ds)
    write_formatted_dataset(ds, "Global_X_Market_Formatted_Dataset.xlsx")

Requires: openpyxl (read), xlsxwriter (write).
"""
from __future__ import annotations

import argparse
import math
import os
import re
import sys
from collections import OrderedDict, defaultdict

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  pip install openpyxl")
try:
    import xlsxwriter
except ImportError:  # pragma: no cover
    sys.exit("xlsxwriter is required:  pip install xlsxwriter")

# --------------------------------------------------------------------------------------------
# Configuration (edit here to change house style)
# --------------------------------------------------------------------------------------------
VALUE, VOLUME = "Value", "Volume"
METRICS = (VALUE, VOLUME)
PLACEHOLDER = "XX"                 # any non-numeric cell (XX, NA, -) propagates as this in sums
CROSS_SEP = " × "                  # segment name for cross segments: "Product Type × Flavour"

GEO_SEGMENT_ALIASES = {"region", "regions", "country", "countries", "region / country",
                       "country / region", "region/country", "geography", "by region",
                       "by country"}
REGION_SEGMENT_NAME = "Region"     # breakdown under the top geography (Global)
COUNTRY_SEGMENT_NAME = "Country"   # breakdown under each region
TOTAL_SEGMENT_ALIASES = {"total market", "total", "overall", "market total"}
OVERALL_LABELS = {"overall", "total", ""}

TITLE_TEMPLATE = "{geo} {market} Market {metrics} Analysis and Forecast, {y0}-{y1}"

BRAND_BLUE = "#0079B2"
ROW_FILL = "#E8E8E8"
BUTTON_FILL = "#B8D5F0"
FONT = "Aptos Narrow"
MASTER_NUM_FMT = '_(* #,##0_);_(* \\(#,##0\\);_(* "-"??_);_(@_)'
PIVOT_NUM_FMT = {VALUE: "#,##0.00", VOLUME: "#,##0"}
MASTER_WIDTHS = [20.9, 20.8, 42.7, 53.1, 53.1, 9.0]   # A..F, then years
MASTER_YEAR_WIDTH = 12.0
PIVOT_LABEL_WIDTH = 65.1
PIVOT_YEAR_WIDTH = 12.8
ZOOM = 70


# --------------------------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------------------------
_TREE_CHARS = re.compile(r"^[\s └├│─\-–•·>]*")


def clean_label(s) -> str:
    """Strip tree glyphs (└ ├ ─), bullets and stray whitespace from a label."""
    if s is None:
        return ""
    s = str(s).replace(" ", " ")
    s = _TREE_CHARS.sub("", s)
    return re.sub(r"\s+", " ", s).strip()


def indent_depth(s) -> int:
    """'X' -> 0, '  └ X' -> 1, '      └ X' -> 2 (4 spaces per extra level)."""
    if s is None:
        return 0
    raw = str(s).replace(" ", " ")
    m = re.match(r"^(\s*)[└├]", raw)
    if m:
        return len(m.group(1)) // 4 + 1
    lead = len(raw) - len(raw.lstrip(" "))
    return lead // 4 if lead >= 4 else 0


def to_num(v):
    """Number -> float; blank -> None; any other text (XX, NA, -) -> PLACEHOLDER."""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v) if math.isfinite(v) else None
    s = str(v).strip()
    if s == "":
        return None
    try:
        return float(s.replace(",", ""))
    except ValueError:
        return PLACEHOLDER


def is_year(v) -> bool:
    try:
        y = int(str(v).strip())
    except (TypeError, ValueError):
        return False
    return 1990 <= y <= 2100 and re.fullmatch(r"\s*\d{4}(\.0)?\s*", str(v)) is not None


_CURRENCY = re.compile(r"(US\$|USD|\$|EUR|€|GBP|£|JPY|¥|INR|₹|CNY|RMB)", re.I)
_VOLUME_WORDS = re.compile(r"\b(units?|MT|tons?|tonnes?|kg|kilo|litres?|liters?|L|pieces?|pcs|"
                           r"doses?|procedures?|tests?|devices?|sq\.? ?m|m2|m3|barrels?|"
                           r"vehicles?|patients?|kits?|packs?|bottles?|cartons?|yards?|yds?|metres?|meters?)\b", re.I)


def classify_unit(label, unit_map=None) -> str:
    """Decide whether a Unit / Metric label is Value or Volume."""
    s = str(label or "").strip()
    if unit_map and s in unit_map:
        return unit_map[s]
    if re.match(r"^value\b", s, re.I):
        return VALUE
    if re.match(r"^volume\b", s, re.I):
        return VOLUME
    if _CURRENCY.search(s) or re.search(r"\b(revenue|sales value|market value)\b", s, re.I):
        return VALUE
    if _VOLUME_WORDS.search(s):
        return VOLUME
    if re.search(r"\b(Mn|Bn|Million|Billion|Thousand)\b", s, re.I):
        return VALUE
    return VOLUME


def unit_text(label) -> str | None:
    """'Value (US$ Mn)' -> 'US$ Mn';  'US$ Mn' -> 'US$ Mn';  'Value' -> None."""
    s = str(label or "").strip()
    m = re.search(r"\(([^)]+)\)", s)
    if m:
        return m.group(1).strip()
    if s.casefold() in ("value", "volume", ""):
        return None
    return s


def sort_key(label):
    """Pivot-style A-Z (used inside segments)."""
    return str(label).casefold()


def geo_sort_key(label):
    """Top-level geography list: A-Z with 'Other ...' / 'Rest of ...' last."""
    s = str(label)
    tail = bool(re.match(r"^(other|others|rest of|rest-of|all other|misc)\b", s, re.I))
    return (tail, s.casefold())


def _sum(vectors, n):
    out = []
    for i in range(n):
        acc, seen, txt = 0.0, False, False
        for vec in vectors:
            x = vec[i]
            if x is None:
                continue
            if isinstance(x, str):
                txt = True
                break
            acc += x
            seen = True
        out.append(PLACEHOLDER if txt else (acc if seen else None))
    return out


def _round_vec(vec, dec):
    return [round(x, dec) if isinstance(x, float) else x for x in vec]


# --------------------------------------------------------------------------------------------
# Data container
# --------------------------------------------------------------------------------------------
class FormattedDataset:
    """Leaf-level rows keyed by (geo, segment, sub, sub1, sub2, metric)."""

    def __init__(self, years, market=None):
        self.years = [int(y) for y in years]
        self.market = market
        self.units = {VALUE: None, VOLUME: None}
        self.rows: "OrderedDict[tuple, list]" = OrderedDict()
        self.geo_order: list[str] = []
        self.geo_children: "OrderedDict[str, list[str]]" = OrderedDict()
        self.geo_segments: set[tuple] = set()          # (geo, segment) that are geo breakdowns
        self.cross_segments: set[str] = set()
        self.reported_totals: dict[tuple, list] = {}    # (geo, metric) -> values
        self.notes: list[str] = []
        # cross rows whose A label carried 'Parent | Child': held back and
        # rolled up to the parent (level one) unless the parent has a row
        self._xchild: dict[tuple, list] = {}

    # ---- building -------------------------------------------------------------------------
    def add(self, geo, segment, sub, sub1=None, sub2=None, metric=VALUE, values=(),
            unit=None, is_geo_segment=False, is_cross=False):
        geo, segment = clean_label(geo), clean_label(segment)
        sub = clean_label(sub)
        sub1 = clean_label(sub1)
        sub2 = clean_label(sub2)
        # 'Parent | Child' in one cell: the text before the bar is the
        # sub-segment and the text after it the next level, in the next column
        raw_sub = sub
        if "|" in sub:
            sub, child = (x.strip() for x in sub.split("|", 1))
            if not is_cross:
                if not sub1 or sub1 == raw_sub:
                    sub1 = child
                if not sub2 or sub2 == raw_sub:
                    sub2 = child
        if sub1.casefold() in OVERALL_LABELS:
            sub1 = sub
        if sub2.casefold() in OVERALL_LABELS:
            sub2 = sub1
        if metric not in METRICS:
            metric = classify_unit(metric)
        vals = [to_num(v) for v in values]
        if len(vals) != len(self.years):
            raise ValueError(f"{geo}|{segment}|{sub1}: {len(vals)} values for {len(self.years)} years")
        if is_cross and "|" in raw_sub:
            # cross analysis runs on level one only: a level-two A row is
            # held, and rolled up to its parent when no parent row exists
            key = (geo, segment, sub, sub1, sub2, metric)
            self._xchild[key] = (_sum([self._xchild[key], vals], self.n)
                                 if key in self._xchild else vals)
            return
        if unit and not self.units.get(metric):
            self.units[metric] = unit
        if geo not in self.geo_order:
            self.geo_order.append(geo)
        if is_geo_segment:
            self.geo_segments.add((geo, segment))
            kids = self.geo_children.setdefault(geo, [])
            if sub not in kids:
                kids.append(sub)
        if is_cross:
            self.cross_segments.add(segment)
        key = (geo, segment, sub, sub1, sub2, metric)
        if key in self.rows:
            self.notes.append(f"duplicate row ignored: {' | '.join(key)}")
            return
        self.rows[key] = vals

    def flush_cross_children(self):
        """Roll the held level-two cross rows up to level one: a parent that
        already has its own row keeps it (the children are its split)."""
        pend, self._xchild = self._xchild, {}
        for key, vals in pend.items():
            if key not in self.rows:
                geo, seg, sub, sub1, sub2, metric = key
                self.add(geo, seg, sub, sub1, sub2, metric, vals, is_cross=True)

    # ---- views ----------------------------------------------------------------------------
    @property
    def n(self):
        return len(self.years)

    def aggregates(self):
        return [g for g in self.geo_order if g in self.geo_children]

    def top_geo(self):
        child = {c for kids in self.geo_children.values() for c in kids}
        for g in self.geo_order:
            if g not in child:
                return g
        return self.geo_order[0] if self.geo_order else "Global"

    def index(self):
        """geo -> metric -> segment -> [(sub, sub1, sub2, vals)] in insertion order."""
        self.flush_cross_children()
        idx = OrderedDict()
        for (g, s, a, b, c, m), v in self.rows.items():
            idx.setdefault(g, OrderedDict()).setdefault(m, OrderedDict()).setdefault(s, []).append((a, b, c, v))
        return idx

    def segment_order(self, geo, segs):
        """Geo segment first, plain segments in source order, cross segments last."""
        geo_s = [s for s in segs if (geo, s) in self.geo_segments]
        cross = [s for s in segs if s in self.cross_segments]
        plain = [s for s in segs if s not in geo_s and s not in cross]
        return geo_s + plain + cross


# --------------------------------------------------------------------------------------------
# Reader 1: FMI Excel Report ('Market Data' + 'Cross-sectional')
# --------------------------------------------------------------------------------------------
def _header(rows, must, limit=40):
    for i, row in enumerate(rows[:limit]):
        low = [str(c).strip().casefold() if c is not None else "" for c in row]
        if all(any(cell == m for cell in low) for m in must):
            return i, low
    return None, None


def _year_cols(header_row, start=0):
    cols, years = [], []
    for j, c in enumerate(header_row):
        if j >= start and is_year(c):
            cols.append(j)
            years.append(int(str(c).strip()[:4]))
    return cols, years


def _market_from_title(title, geo=None):
    t = clean_label(title)
    if geo and t.casefold().startswith(geo.casefold() + " "):
        t = t[len(geo) + 1:]
    m = re.match(r"^(.*?)\s+Market\b", t)
    return m.group(1).strip() if m else None


def _leaf_paths(entries):
    """entries: [(depth, name, vals)] -> leaves [(path, vals)] and all node paths {name: path}."""
    out, nodes, stack = [], {}, []
    for i, (d, name, vals) in enumerate(entries):
        stack = stack[:d]
        path = stack + [name]
        stack = path
        nodes.setdefault(name.casefold(), path)
        nxt = entries[i + 1][0] if i + 1 < len(entries) else -1
        if nxt <= d:
            out.append((path, vals))
    return out, nodes


def _path_cols(path):
    if len(path) == 1:
        return path[0], path[0], path[0]
    if len(path) == 2:
        return path[0], path[1], path[1]
    return path[0], path[-2], path[-1]


def read_excel_report(path, market=None, include_cross=True) -> FormattedDataset:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # ---- locate 'Market Data' ----
    md_rows = hi = low = None
    names = sorted(wb.sheetnames, key=lambda n: 0 if "market data" in n.casefold() else 1)
    for name in names:
        rows = list(wb[name].iter_rows(values_only=True))
        h, l = _header(rows, ("geography", "segment", "metric"))
        if h is not None and "segment pair" not in l:
            md_rows, hi, low = rows, h, l
            break
    if md_rows is None:
        raise ValueError("No 'Market Data'-style sheet (Geography | Segment | Sub-segment | Metric | years) found")

    gi, si, mi = low.index("geography"), low.index("segment"), low.index("metric")
    ssi = low.index("sub-segment") if "sub-segment" in low else si + 1
    ycols, years = _year_cols(md_rows[hi])
    title = next((r[0] for r in md_rows[:hi] if r and r[0]), "")

    # pass 1: group rows per (geo, segment, metric) in order
    groups = OrderedDict()
    units = {}
    for r in md_rows[hi + 1:]:
        if not r or r[gi] is None or r[mi] is None or r[si] is None:
            continue
        geo, seg, metric_lbl = clean_label(r[gi]), clean_label(r[si]), str(r[mi])
        metric = classify_unit(metric_lbl)
        units.setdefault(metric, unit_text(metric_lbl))
        vals = [r[j] if j < len(r) else None for j in ycols]
        groups.setdefault((geo, seg, metric), []).append((indent_depth(r[ssi]), clean_label(r[ssi]), vals))

    geo_order = list(OrderedDict.fromkeys(g for g, _, _ in groups))
    aggregates = {g for g, s, _ in groups if s.casefold() in GEO_SEGMENT_ALIASES}
    top = geo_order[0] if geo_order else "Global"
    ds = FormattedDataset(years, market or _market_from_title(title, top) or _market_from_title(title))
    ds.units.update({k: v for k, v in units.items() if v})

    taxonomy = defaultdict(dict)   # (geo, segment) -> {name_cf: path}
    for (geo, seg, metric), entries in groups.items():
        segcf = seg.casefold()
        if segcf in TOTAL_SEGMENT_ALIASES:
            ds.reported_totals[(geo, metric)] = [to_num(v) for v in entries[0][2]]
            if geo not in ds.geo_order:
                ds.geo_order.append(geo)
            continue
        if segcf in GEO_SEGMENT_ALIASES:
            kids = [e for e in entries if e[0] == 0]
            seg_name = REGION_SEGMENT_NAME if any(k[1] in aggregates for k in kids) else COUNTRY_SEGMENT_NAME
            for _, name, vals in kids:
                ds.add(geo, seg_name, name, name, name, metric, vals, is_geo_segment=True)
            continue
        leaves, nodes = _leaf_paths(entries)
        taxonomy[(geo, seg)].update(nodes)
        for p, vals in leaves:
            if len(p) > 3:
                ds.notes.append(f"{seg}: depth {len(p)} path collapsed to 3 levels: {' > '.join(p)}")
            a, b, c = _path_cols(p)
            ds.add(geo, seg, a, b, c, metric, vals)

    # keep source geo order (incl. geos that only had totals)
    ds.geo_order = [g for g in geo_order if g in ds.geo_order] + [g for g in ds.geo_order if g not in geo_order]

    if include_cross:
        _read_cross(wb, ds, taxonomy, top)
    wb.close()
    return ds


def _read_cross(wb, ds, taxonomy, top):
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(values_only=True))
        hi, low = _header(rows, ("geography", "segment pair", "metric"))
        if hi is None:
            continue
        gi, pi, mi = low.index("geography"), low.index("segment pair"), low.index("metric")
        d1i = low.index("dimension 1") if "dimension 1" in low else None
        d2i = low.index("dimension 2") if "dimension 2" in low else None
        ycols, years = _year_cols(rows[hi])
        if years != ds.years:
            ds.notes.append(f"'{name}': year columns differ from Market Data, sheet skipped")
            continue
        groups = OrderedDict()
        for r in rows[hi + 1:]:
            if not r or r[gi] is None or r[pi] is None or r[mi] is None:
                continue
            pair = str(r[pi])
            parts = re.split(r"\s*×\s*", pair) if "×" in pair else re.split(r"\s+x\s+", pair, flags=re.I)
            if len(parts) != 2:
                ds.notes.append(f"cross pair not understood, skipped: {pair}")
                continue
            a, b = clean_label(parts[0]), clean_label(parts[1])
            d1 = clean_label(r[d1i]) if d1i is not None else "Dimension 1"
            d2 = clean_label(r[d2i]) if d2i is not None else "Dimension 2"
            metric = classify_unit(r[mi])
            vals = [r[j] if j < len(r) else None for j in ycols]
            groups.setdefault((clean_label(r[gi]), d1, d2, metric), []).append((a, b, vals))

        def tax(geo, dim):
            return taxonomy.get((geo, dim)) or taxonomy.get((top, dim)) or {}

        for (geo, d1, d2, metric), items in groups.items():
            t1, t2 = tax(geo, d1), tax(geo, d2)
            As = {a.casefold() for a, _, _ in items}
            Bs = {b.casefold() for _, b, _ in items}

            def has_child(name, pool, t):
                cf = name.casefold()
                return any(cf in [x.casefold() for x in t.get(o, [])[:-1]] for o in pool if o != cf)

            seg = f"{d1}{CROSS_SEP}{d2}"
            for a, b, vals in items:
                if has_child(a, As, t1) or has_child(b, Bs, t2):
                    continue                       # parent pair: children carry the split
                pa = t1.get(a.casefold(), [a])
                ds.add(geo, seg, pa[0], a, b, metric, vals, is_cross=True)


# --------------------------------------------------------------------------------------------
# Reader 2: FMI Raw Data workbook (or an existing Formatted Dataset 'Master Sheet')
# --------------------------------------------------------------------------------------------
def read_raw_data(path, market=None, sheet=None, unit_map=None) -> FormattedDataset:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    order = [sheet] if sheet else sorted(
        wb.sheetnames, key=lambda n: (0 if n.casefold() in ("raw data", "master sheet") else 1))
    for name in order:
        rows = list(wb[name].iter_rows(values_only=True))
        hi, low = _header(rows, ("region", "segment", "unit"))
        if hi is not None:
            break
    else:
        raise ValueError("No Raw Data-style sheet (Region | Segment | ... | Unit | years) found")

    ri, si, ui = low.index("region"), low.index("segment"), low.index("unit")
    s0 = low.index("sub-segment") if "sub-segment" in low else si + 1
    s1 = low.index("sub-segment 1") if "sub-segment 1" in low else None
    s2 = low.index("sub-segment 2") if "sub-segment 2" in low else None
    ycols, years = _year_cols(rows[hi], start=ui + 1)
    # stop at the first gap (H-CAGR / YoY blocks follow a blank column)
    cut = [ycols[0]]
    for j in ycols[1:]:
        if j != cut[-1] + 1:
            break
        cut.append(j)
    ycols, years = cut, years[:len(cut)]
    title = next((r[0] for r in rows[:hi] if r and r[0]), "")

    ds = FormattedDataset(years, market)
    # units from the title, e.g. "... Value (US$ Mn) and Volume (MT) ..."
    for metric in METRICS:
        m = re.search(metric + r"\s*\(([^)]+)\)", str(title), re.I)
        if m:
            ds.units[metric] = m.group(1).strip()

    pending = []
    for r in rows[hi + 1:]:
        if not r or r[ri] is None or r[si] is None or r[ui] is None:
            continue
        unit_lbl = str(r[ui]).strip()
        metric = classify_unit(unit_lbl, unit_map)
        u = unit_text(unit_lbl)
        vals = [r[j] if j < len(r) else None for j in ycols]
        pending.append((r[ri], r[si], r[s0], r[s1] if s1 is not None else None,
                        r[s2] if s2 is not None else None, metric, vals, u))
    agg_names = {clean_label(p[0]) for p in pending if clean_label(p[1]).casefold() in GEO_SEGMENT_ALIASES}
    for geo, seg, a, b, c, metric, vals, u in pending:
        segcf = clean_label(seg).casefold()
        is_geo = segcf in GEO_SEGMENT_ALIASES
        if is_geo:
            seg = REGION_SEGMENT_NAME if clean_label(a) in agg_names else COUNTRY_SEGMENT_NAME
        ds.add(geo, seg, a, b, c, metric, vals, unit=u, is_geo_segment=is_geo,
               is_cross=(CROSS_SEP.strip() in str(seg)))
    if not ds.market:
        ds.market = _market_from_title(title, ds.top_geo()) or _market_from_title(title)
    wb.close()
    return ds


# --------------------------------------------------------------------------------------------
# Reader 3: in-memory records  /  Reader 4: segmentation only
# --------------------------------------------------------------------------------------------
_REC_KEYS = {
    "geo": ("geo", "geography", "region"),
    "segment": ("segment",),
    "sub": ("sub", "sub-segment", "sub_segment", "subsegment"),
    "sub1": ("sub1", "sub-segment 1", "sub_segment_1"),
    "sub2": ("sub2", "sub-segment 2", "sub_segment_2"),
    "metric": ("metric", "unit"),
}


def from_records(records, years, market, value_unit="US$ Mn", volume_unit="MT") -> FormattedDataset:
    """records: iterable of dicts with geo/segment/sub/sub1/sub2/metric and either a
    'values' list or one key per year (2021 or '2021'). Keys are case-insensitive."""
    ds = FormattedDataset(years, market)
    ds.units = {VALUE: value_unit, VOLUME: volume_unit}
    for rec in records:
        low = {str(k).strip().casefold(): v for k, v in rec.items()}
        get = lambda f: next((low[k] for k in _REC_KEYS[f] if k in low), None)  # noqa: E731
        vals = low.get("values") or [low.get(str(y), rec.get(y)) for y in ds.years]
        seg = clean_label(get("segment"))
        is_geo = seg.casefold() in GEO_SEGMENT_ALIASES
        ds.add(get("geo"), seg, get("sub"), get("sub1"), get("sub2"),
               classify_unit(get("metric")), vals, is_geo_segment=is_geo,
               is_cross=CROSS_SEP.strip() in seg)
    return ds


def _walk(node, prefix=()):
    if isinstance(node, dict):
        for k, v in node.items():
            if v in (None, [], {}):
                yield prefix + (k,)
            else:
                yield from _walk(v, prefix + (k,))
    else:
        for leaf in node:
            if isinstance(leaf, dict):
                yield from _walk(leaf, prefix)
            else:
                yield prefix + (leaf,)


# The standard FMI regional taxonomy: used whenever neither the segmentation,
# the raw data nor the instruction names any region or country.
STANDARD_GEOGRAPHIES = OrderedDict([
    ("North America", ["U.S.", "Canada"]),
    ("Latin America", ["Brazil", "Mexico", "Colombia", "Rest of Latin America"]),
    ("Europe", ["Germany", "U.K.", "France", "Italy", "Spain", "Belgium",
                "Netherlands", "Nordic Countries", "Switzerland", "Russia",
                "Poland", "Rest of Europe"]),
    ("East Asia", ["China", "Japan", "South Korea"]),
    ("South Asia & Pacific", ["India", "Thailand", "Malaysia", "Indonesia",
                              "Singapore", "Philippines", "Australia",
                              "New Zealand", "Rest of South Asia & Pacific"]),
    ("Middle East & Africa", ["Turkey", "UAE", "Saudi Arabia", "Israel",
                              "Other GCC Countries", "North Africa",
                              "South Africa", "Rest of Middle East & Africa"]),
])


def align_to_top_geography(ds) -> FormattedDataset:
    """The top geography's segmentation is the authority: a leaf that a
    region or country carries but the top geography does not (rows left over
    from another report, a stray extra sub-segment) cannot be summed into it,
    so it is dropped - and noted - rather than failing the tie-out."""
    if not ds.geo_children or not ds.rows:
        return ds
    top = ds.top_geo()
    allowed = {}
    for (g, seg, a, b, c, m) in ds.rows:
        if g == top:
            allowed.setdefault(seg, set()).add((a, b, c))
    dropped = []
    for key in list(ds.rows):
        g, seg, a, b, c, m = key
        if g == top or (g, seg) in ds.geo_segments or seg in ds.cross_segments:
            continue
        if seg in allowed and (a, b, c) not in allowed[seg]:
            dropped.append(key)
            del ds.rows[key]
    if dropped:
        ds.notes.append(
            f"{len(dropped)} rows dropped: leaves a region / country carries "
            f"that {top} does not (e.g. {' | '.join(dropped[0][:3])})")
    return ds


def ensure_standard_geographies(ds, geographies=None) -> FormattedDataset:
    """Give a dataset that carries no region / country the standard regional
    taxonomy.  A dataset that already names any region or country is left
    exactly as it is.

    Every segment row of the top geography is repeated for each region and
    country as a placeholder row (the figures are for the research team, or
    for fmi_indicative_numbers.py, to fill); the top geography keeps what it
    has, and the Region / Country breakdown rows are added.
    """
    if ds.geo_children or not ds.rows:
        return ds
    geographies = geographies or STANDARD_GEOGRAPHIES
    top = ds.top_geo()
    blank = [PLACEHOLDER] * len(ds.years)
    base = [(k, v) for k, v in ds.rows.items() if k[0] == top]
    if not base:
        return ds
    geos = list(geographies) + [c for cs in geographies.values() for c in cs]
    for g in [top] + geos:
        for m in METRICS:
            if g == top:
                for rg in geographies:
                    ds.add(g, REGION_SEGMENT_NAME, rg, rg, rg, m, blank,
                           is_geo_segment=True)
            elif g in geographies:
                for c in geographies[g]:
                    ds.add(g, COUNTRY_SEGMENT_NAME, c, c, c, m, blank,
                           is_geo_segment=True)
        if g == top:
            continue
        for (_g, seg, a, b, c, m), _v in base:
            ds.add(g, seg, a, b, c, m, blank, is_cross=seg in ds.cross_segments)
    ds.notes.append("no region / country in the source: standard FMI regional "
                    f"taxonomy applied ({len(geographies)} regions, "
                    f"{sum(len(v) for v in geographies.values())} countries)")
    return ds


def from_taxonomy(taxonomy, geographies, years, market, top_geo="Global", cross=(),
                  fill=PLACEHOLDER, value_unit="US$ Mn", volume_unit="MT") -> FormattedDataset:
    """Skeleton dataset from a segmentation.

    taxonomy    {"Product Type": {"Sweet muffins": ["Fruit muffins", ...], "Savoury muffins": [...]},
                 "Storage Format": ["Ambient", "Chilled", "Frozen"], ...}
    geographies {"North America": ["U.S.", "Canada"], "Europe": [...], ...}
    cross       [("Product Type", "Flavour"), ...]  -> leaves of dim 1 × leaves of dim 2
    """
    ds = FormattedDataset(years, market)
    ds.units = {VALUE: value_unit, VOLUME: volume_unit}
    if not geographies:                       # none given: the standard set
        geographies = STANDARD_GEOGRAPHIES
    blank = [fill] * len(ds.years)
    leaves = {seg: list(_walk(tree)) for seg, tree in taxonomy.items()}
    geos = [top_geo] + list(geographies) + [c for cs in geographies.values() for c in cs]
    for g in geos:
        for m in METRICS:
            if g == top_geo:
                for rg in geographies:
                    ds.add(g, REGION_SEGMENT_NAME, rg, rg, rg, m, blank, is_geo_segment=True)
            elif g in geographies:
                for c in geographies[g]:
                    ds.add(g, COUNTRY_SEGMENT_NAME, c, c, c, m, blank, is_geo_segment=True)
            for seg, paths in leaves.items():
                for p in paths:
                    a, b, c = _path_cols(list(p))
                    ds.add(g, seg, a, b, c, m, blank)
            for d1, d2 in cross:
                for pa in leaves[d1]:
                    for pb in leaves[d2]:
                        ds.add(g, f"{d1}{CROSS_SEP}{d2}", pa[0], pa[-1], pb[-1], m, blank, is_cross=True)
    return ds


# Reader 5: a SEGMENTATION workbook - one sheet, "Parameter | Value |
# Segmentation | Companies | ..." - the basic brief the research team can hand
# over instead of a dataset (headline figures + taxonomy + companies).
_AGE_WORDS = {"paediatric", "pediatric", "adult", "geriatric", "elderly",
              "infant", "infants", "children", "senior", "seniors", "adults"}


def _dim_name(label):
    t = clean_label(label)
    return re.sub(r"^by\s+", "", t, flags=re.I).strip() or t


def _num(v):
    try:
        return float(str(v).replace(",", "").replace("%", "").strip())
    except (TypeError, ValueError):
        return None


def is_segmentation(path) -> bool:
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=5, values_only=True):
                low = [str(c).strip().lower() for c in row if c is not None]
                if "segmentation" in low and "parameter" in low:
                    wb.close()
                    return True
        wb.close()
    except Exception:
        pass
    return False


def read_segmentation(path, market=None, years=None, keep_cells=False):
    """Skeleton ('XX') dataset from a segmentation workbook.

    Reads, from the first sheet that has a Parameter and a Segmentation column:
      * headline parameters (2026 / 2036 Value, CAGR, Volume unit) -> ds.headline
      * the taxonomy: bold cells are the dimensions ("By Product Type"), the
        cell's indent is the level beneath it; a dimension named twice gets a
        second name ("Age Group" when its members are ages, else "(2)")
      * the "Region" block when present (Global > regions > countries, by
        indent) - otherwise the standard regional taxonomy
      * the Companies column -> ds.companies
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = None
    hdr = None
    for w in wb.worksheets:
        for r in range(1, min(w.max_row, 6) + 1):
            low = {str(w.cell(r, c).value or "").strip().lower(): c
                   for c in range(1, w.max_column + 1)}
            if "segmentation" in low and "parameter" in low:
                ws, hdr = w, (r, low)
                break
        if ws:
            break
    if ws is None:
        raise ValueError(f"{path}: no Parameter | Segmentation sheet found")
    r0, low = hdr
    cp, cv, cs = low["parameter"], low.get("value", low["parameter"] + 1),         low["segmentation"]
    cc = low.get("companies")

    head = {"value_2025": None, "value_2026": None, "value_2036": None,
            "cagr": None, "volume_unit": None}
    companies = []
    dims = []            # [(name, [(indent, text)])]
    region_items = None
    cur = None
    for r in range(r0 + 1, ws.max_row + 1):
        p = ws.cell(r, cp).value
        if p is not None:
            pl = str(p).strip().lower()
            v = ws.cell(r, cv).value
            m = re.match(r"(\d{4})\s*value", pl)
            if m and _num(v) is not None and "value_" + m.group(1) in head:
                head["value_" + m.group(1)] = _num(v)
            elif pl.startswith("cagr") and _num(v) is not None:
                head["cagr"] = _num(v)
            elif pl.startswith("volume") and v:
                head["volume_unit"] = str(v).strip()
        if cc:
            c = ws.cell(r, cc).value
            if c is not None and str(c).strip():
                companies.append(clean_label(c))
        cell = ws.cell(r, cs)
        t = cell.value
        if t is None or not str(t).strip():
            continue
        t = clean_label(t)
        indent = int(cell.alignment.indent or 0)
        if cell.font.b or (indent == 0 and t.lower().startswith("by ")):
            cur = [t, []]
            dims.append(cur)
        elif cur is not None:
            cur[1].append((indent, t))

    taxonomy = OrderedDict()
    geographies = None
    for name, items in dims:
        if name.strip().lower() == "region":
            region_items = items
            continue
        dn = _dim_name(name)
        if dn in taxonomy:
            members = {i[1].strip().lower() for i in items}
            alt = "Age Group" if members & _AGE_WORDS else f"{dn} (2)"
            dn = alt if alt not in taxonomy else f"{alt} (2)"
        base = min((i for i, _ in items), default=1)
        tree = OrderedDict()
        flat = True
        last = None
        for ind, t in items:
            if ind <= base or last is None:
                tree[t] = []
                last = t
            else:
                tree[last].append(t)
                flat = False
        taxonomy[dn] = (list(tree) if flat
                        else OrderedDict((k, v) for k, v in tree.items()))

    if region_items:
        regs = OrderedDict()
        lvls = sorted({i for i, _ in region_items})
        top_ind = lvls[0]
        # the first item is the top geography ("Global"); regions sit one
        # level in, countries two - a region named again simply continues it
        reg_ind = lvls[1] if len(lvls) > 1 else None
        cur_reg = None
        for ind, t in region_items:
            if ind == top_ind or reg_ind is None:
                continue
            if ind == reg_ind:
                cur_reg = t
                regs.setdefault(t, [])
            elif cur_reg is not None:
                if t not in regs[cur_reg]:
                    regs[cur_reg].append(t)
        if regs and any(regs.values()):
            geographies = regs

    if market is None:
        market = os.path.splitext(os.path.basename(str(path)))[0]
        market = re.sub(r"[_ ]*segmentation.*$", "", market, flags=re.I)
        market = re.sub(r"\s+market$", "", market.replace("_", " "),
                        flags=re.I).strip()
    ds = from_taxonomy(taxonomy, geographies, years or list(range(2021, 2037)),
                       market, volume_unit=head["volume_unit"] or "MT")
    ds.headline = head
    ds.companies = companies
    ds.notes.append(
        f"built from a segmentation: {len(taxonomy)} dimensions"
        + ("" if geographies else "; no regions given - standard regional "
                                  "taxonomy used"))
    return ds


# --------------------------------------------------------------------------------------------
# Reader 5: the research team's 'Market Data' + 'Cross-sectional' layout
#   Market Size | Region/ Country | Segment Type | Sub Segment Type | years
#   Market Size | Region/ Country | <A axis> | Cross Segment | Sub-Segment | years
# 'Parent | Child' in a sub-segment cell is read as sub-segment | next level.
# --------------------------------------------------------------------------------------------
def _is_market_data_layout(rows):
    h, _ = _header(rows, ("market size", "segment type", "sub segment type"))
    return h is not None


def read_market_data_layout(path, market=None) -> FormattedDataset:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {n: list(wb[n].iter_rows(values_only=True)) for n in wb.sheetnames}
    wb.close()
    main = next((n for n, r in sheets.items() if _is_market_data_layout(r)), None)
    if main is None:
        raise ValueError("No 'Market Data' sheet (Market Size | Region/ Country | Segment Type | "
                         "Sub Segment Type) found")
    rows = sheets[main]
    hi, low = _header(rows, ("market size", "segment type", "sub segment type"))
    mi, si, ssi = low.index("market size"), low.index("segment type"), low.index("sub segment type")
    ri = next(j for j, c in enumerate(low) if c.startswith("region"))
    ycols, years = _year_cols(rows[hi])
    title = next((r[0] for r in rows[:hi]
                  if r and r[0] and not str(r[0]).lower().startswith("note")), "")
    ds = FormattedDataset(years, market or _market_from_title(title))

    def val(r, j):
        return r[j] if j < len(r) else None

    body = [r for r in rows[hi + 1:]
            if r and r[mi] is not None and r[ri] is not None and r[si] is not None]
    geos = {clean_label(r[ri]) for r in body}
    for r in body:
        metric = classify_unit(r[mi])
        u = unit_text(r[mi])
        seg = re.sub(r"^by\s+", "", clean_label(r[si]), flags=re.I)
        sub = clean_label(r[ssi])
        vals = [val(r, j) for j in ycols]
        if seg.casefold() in GEO_SEGMENT_ALIASES:
            seg = REGION_SEGMENT_NAME if sub in geos else COUNTRY_SEGMENT_NAME
            ds.add(r[ri], seg, sub, sub, sub, metric, vals, unit=u, is_geo_segment=True)
        else:
            ds.add(r[ri], seg, sub, None, None, metric, vals, unit=u)

    for n, rws in sheets.items():
        h, lw = _header(rws, ("market size", "cross segment", "sub-segment"))
        if h is None:
            continue
        m_i, c_i, b_i = lw.index("market size"), lw.index("cross segment"), lw.index("sub-segment")
        r_i = next(j for j, c in enumerate(lw) if c.startswith("region"))
        a_i = r_i + 1
        yc, yrs = _year_cols(rws[h])
        if yrs != ds.years:
            ds.notes.append(f"'{n}': year columns differ from Market Data, sheet skipped")
            continue
        for r in rws[h + 1:]:
            if not r or r[m_i] is None or r[r_i] is None or r[c_i] is None:
                continue
            parts = re.split(r"\s*[\u00d7\ufffd]\s*|\s+x\s+", clean_label(r[c_i]), maxsplit=1,
                             flags=re.I)
            seg = (f"{parts[0].strip()}{CROSS_SEP}{parts[1].strip()}" if len(parts) == 2
                   else clean_label(r[c_i]))
            a, b = clean_label(val(r, a_i)), clean_label(val(r, b_i))
            ds.add(r[r_i], seg, a, b, b, classify_unit(r[m_i]), [val(r, j) for j in yc],
                   unit=unit_text(r[m_i]), is_cross=True)
    ds.flush_cross_children()
    return ds


def load_any(path, market=None, **kw) -> FormattedDataset:
    """Auto-detect: Excel Report (Geography/Metric) or Raw Data (Region/Unit)
    or a Segmentation brief (Parameter/Segmentation)."""
    if is_segmentation(path):
        return read_segmentation(path, market=market)
    wb = openpyxl.load_workbook(path, read_only=True)
    kind = None
    for name in wb.sheetnames:
        rows = list(wb[name].iter_rows(max_row=40, values_only=True))
        if _is_market_data_layout(rows):
            kind = "marketdata"
            break
        if _header(rows, ("geography", "segment", "metric"))[0] is not None:
            kind = "report"
            break
        if _header(rows, ("region", "segment", "unit"))[0] is not None:
            kind = "raw"
    wb.close()
    if kind == "marketdata":
        return align_to_top_geography(ensure_standard_geographies(
            read_market_data_layout(path, market=market)))
    if kind == "report":
        return align_to_top_geography(ensure_standard_geographies(
            read_excel_report(path, market=market)))
    if kind == "raw":
        return align_to_top_geography(ensure_standard_geographies(
            read_raw_data(path, market=market, **kw)))
    raise ValueError(f"{path}: neither an FMI Excel Report, a Market Data layout nor a Raw Data workbook")


# --------------------------------------------------------------------------------------------
# QC
# --------------------------------------------------------------------------------------------
def run_qc(ds: FormattedDataset, value_decimals=3, volume_decimals=3, rel_tol=1e-3, error_tol=5e-3):
    """Returns [(level, message)], level in {'ERROR', 'WARN', 'INFO'}. Numeric checks only run
    where every term is numeric; placeholder (XX) data is reported, not failed."""
    issues = []
    idx = ds.index()
    n = ds.n
    dec = {VALUE: value_decimals, VOLUME: volume_decimals}

    def close(a, b, terms, m):
        if not (isinstance(a, float) and isinstance(b, float)):
            return True
        return abs(a - b) <= max(rel_tol * abs(b), terms * 0.5 * 10 ** -dec[m] + 1e-9)

    def first_bad(x, y, terms, m):
        for i in range(n):
            if not close(x[i], y[i], terms, m):
                return i
        return None

    def sev(x, y):
        """ERROR above error_tol relative gap, WARN for small drift (rounding in the source)."""
        return "ERROR" if abs(x - y) > error_tol * max(abs(y), 1e-9) else "WARN"

    totals = {}
    for g, by_m in idx.items():
        for m, by_s in by_m.items():
            segs = ds.segment_order(g, list(by_s))
            sums = {s: _sum([r[3] for r in by_s[s]], n) for s in segs}
            ref_s = segs[0]
            totals[(g, m)] = sums[ref_s]
            for s in segs[1:]:
                i = first_bad(sums[s], sums[ref_s], len(by_s[s]), m)
                if i is not None:
                    issues.append((sev(sums[s][i], sums[ref_s][i]), f"{g} | {m} | '{s}' sums to {sums[s][i]:,.3f} vs '{ref_s}' "
                                            f"{sums[ref_s][i]:,.3f} in {ds.years[i]}"))
            rep = ds.reported_totals.get((g, m))
            if rep:
                i = first_bad(sums[ref_s], rep, len(by_s[ref_s]), m)
                if i is not None:
                    issues.append((sev(sums[ref_s][i], rep[i]), f"{g} | {m} | leaves {sums[ref_s][i]:,.3f} vs reported "
                                            f"Total Market {rep[i]:,.3f} in {ds.years[i]}"))
            for s in segs:
                for a, b, c, v in by_s[s]:
                    if any(isinstance(x, float) and x < 0 for x in v):
                        issues.append(("ERROR", f"negative value: {g} | {s} | {b} | {c} | {m}"))

    # aggregate geography == sum of its children (totals and every segment leaf)
    for g, kids in ds.geo_children.items():
        for m in METRICS:
            if (g, m) not in totals:
                continue
            kid_tot = [totals[(k, m)] for k in kids if (k, m) in totals]
            if len(kid_tot) == len(kids) and kids:
                s = _sum(kid_tot, n)
                i = first_bad(totals[(g, m)], s, len(kids), m)
                if i is not None:
                    issues.append((sev(totals[(g, m)][i], s[i]), f"{g} | {m} total {totals[(g, m)][i]:,.3f} != Σ children "
                                            f"{s[i]:,.3f} in {ds.years[i]}"))
            if not all(k in idx and m in idx[k] for k in kids):
                continue
            for seg, rows in idx[g][m].items():
                if (g, seg) in ds.geo_segments:
                    continue
                kid_maps = []
                for k in kids:
                    kr = idx[k][m].get(seg)
                    if not kr:
                        break
                    kid_maps.append({(a, b, c): v for a, b, c, v in kr})
                else:
                    bad, worst = 0, 0.0
                    for a, b, c, v in rows:
                        vecs = [km.get((a, b, c)) for km in kid_maps]
                        if any(x is None for x in vecs):
                            continue
                        s = _sum(vecs, n)
                        i = first_bad(v, s, len(kids), m)
                        if i is not None:
                            bad += 1
                            worst = max(worst, max(abs(v[k] - s[k]) / max(abs(s[k]), 1e-9) for k in range(n)
                                                   if isinstance(v[k], float) and isinstance(s[k], float)))
                    if bad:
                        issues.append(("ERROR" if worst > error_tol else "WARN",
                                       f"{g} | {m} | '{seg}': {bad} leaf rows != Σ {len(kids)} children "
                                       f"(worst gap {worst:.2%})"))

    # Value / Volume pairing
    keys = defaultdict(set)
    for (g, s, a, b, c, m) in ds.rows:
        keys[(g, s, a, b, c)].add(m)
    lone = [k for k, ms in keys.items() if len(ms) == 1]
    if lone:
        issues.append(("WARN", f"{len(lone)} leaf rows have only one metric, e.g. {' | '.join(lone[0])}"))

    ph = sum(1 for v in ds.rows.values() for x in v if x == PLACEHOLDER)
    if ph:
        issues.append(("INFO", f"{ph:,} placeholder ({PLACEHOLDER}) cells carried through; numeric ties skipped there"))
    for note in ds.notes[:50]:
        issues.append(("WARN", note))
    if len(ds.notes) > 50:
        issues.append(("WARN", f"... {len(ds.notes) - 50} more notes"))
    return issues


# --------------------------------------------------------------------------------------------
# Writer
# --------------------------------------------------------------------------------------------
def _pivot_rows(ds, metric, dec):
    """[(level, label, vals)] for one metric, plus the Grand Total vector."""
    idx = ds.index()
    n = ds.n
    aggs = [g for g in ds.aggregates() if g in idx]
    others = sorted([g for g in ds.geo_order if g in idx and g not in aggs], key=geo_sort_key)
    out, totals = [], {}
    for g in aggs + others:
        by_s = idx[g].get(metric)
        if not by_s:
            continue
        rounded = {s: [(a, b, c, _round_vec(v, dec)) for a, b, c, v in rows] for s, rows in by_s.items()}
        ref = ds.segment_order(g, list(rounded))[0]
        gtot = _sum([r[3] for r in rounded[ref]], n)
        totals[g] = gtot
        out.append((0, g, gtot))
        for s in sorted(rounded, key=sort_key):
            rows = rounded[s]
            out.append((1, s, _sum([r[3] for r in rows], n)))
            tree = OrderedDict()
            for a, b, c, v in rows:
                tree.setdefault(a, OrderedDict()).setdefault(b, []).append((c, v))
            for a in sorted(tree, key=sort_key):
                out.append((2, a, _sum([v for lst in tree[a].values() for _, v in lst], n)))
                for b in sorted(tree[a], key=sort_key):
                    leaves = tree[a][b]
                    out.append((3, b, _sum([v for _, v in leaves], n)))
                    if any(c != b for c, _ in leaves):
                        merged = OrderedDict()
                        for c, v in leaves:
                            merged.setdefault(c, []).append(v)
                        for c in sorted(merged, key=sort_key):
                            out.append((4, c, _sum(merged[c], n)))
    top = ds.top_geo()
    grand = totals.get(top) or (next(iter(totals.values())) if totals else [None] * n)
    return out, grand


def _title(ds, metrics, top):
    def part(m):
        u = ds.units.get(m)
        return f"{m} ({u})" if u else m
    return TITLE_TEMPLATE.format(geo=top, market=ds.market or "", y0=ds.years[0], y1=ds.years[-1],
                                 metrics=" and ".join(part(m) for m in metrics)).replace("  ", " ")


def write_formatted_dataset(ds: FormattedDataset, out_path, unit_style="metric",
                            value_decimals=3, volume_decimals=3, buttons=True):
    """unit_style='metric' -> Unit column says Value / Volume (default, FMI formatted dataset)
       unit_style='unit'   -> Unit column says US$ Mn / MT (the ds.units labels)."""
    n = ds.n
    dec = {VALUE: value_decimals, VOLUME: volume_decimals}
    top = ds.top_geo()
    present = [m for m in METRICS if any(k[5] == m for k in ds.rows)]
    unit_lbl = {m: (ds.units.get(m) or m) if unit_style == "unit" else m for m in METRICS}
    last_col = 6 + n - 1

    wb = xlsxwriter.Workbook(out_path, {"constant_memory": True, "strings_to_numbers": False})
    F = lambda **k: wb.add_format({"font_name": FONT, "font_size": 11, **k})  # noqa: E731
    f_title = F(bold=True, font_size=18, font_color="#FFFFFF", bg_color=BRAND_BLUE, align="center", valign="vcenter")
    f_h_txt = F(bold=True, font_color="#FFFFFF", bg_color=BRAND_BLUE, align="left", valign="vcenter", border=1)
    f_h_yr = F(bold=True, font_color="#FFFFFF", bg_color=BRAND_BLUE, align="center", valign="vcenter", border=1)
    f_txt = F(bg_color=ROW_FILL, border=1)
    f_num = F(bg_color=ROW_FILL, border=1, align="center", num_format=MASTER_NUM_FMT)
    f_plain = F()
    f_lbl = [F(align="left", indent=i) for i in range(6)]
    f_pnum = {m: F(align="center", valign="vcenter", num_format=PIVOT_NUM_FMT[m]) for m in METRICS}
    btn_font = {"name": FONT, "size": 18, "color": "#000000"}
    btn_align = {"vertical": "middle", "horizontal": "center"}

    def button(ws, text, target, row, col, x_off=0, y_off=0, w=192, h=77, dashed=True):
        if not buttons:
            return
        ws.insert_textbox(row, col, text, {
            "width": w, "height": h, "x_offset": x_off, "y_offset": y_off,
            "font": btn_font, "align": btn_align, "fill": {"color": BUTTON_FILL},
            "line": {"color": "#000000", "dash_type": "dash"} if dashed else {"none": True},
            "url": f"internal:{target}!A1", "tip": f"Click here to go to {text}",
            "object_position": 3})

    ref = {"Home": "Home", "Master Sheet": "'Master Sheet'", VALUE: VALUE, VOLUME: VOLUME}

    # ---- Home ----
    home = wb.add_worksheet("Home")
    home.hide_gridlines(2)
    home.activate()
    targets = [(f"{VALUE} Data", ref[VALUE]) if VALUE in present else None,
               ("Master Sheet", ref["Master Sheet"]),
               (f"{VOLUME} Data", ref[VOLUME]) if VOLUME in present else None]
    for i, t in enumerate([t for t in targets if t]):
        button(home, t[0], t[1], 6 + i * 4, 14, w=178, h=57, dashed=False)

    # ---- Master Sheet ----
    ms = wb.add_worksheet("Master Sheet")
    ms.hide_gridlines(2)
    ms.set_zoom(ZOOM)
    for j, w in enumerate(MASTER_WIDTHS):
        ms.set_column(j, j, w)
    ms.set_column(6, last_col, MASTER_YEAR_WIDTH)
    ms.merge_range(0, 0, 3, last_col, _title(ds, present, top), f_title)
    button(ms, "Home", ref["Home"], 0, 1, y_off=4)
    others = [m for m in present]
    for k, m in enumerate(others):
        button(ms, m, ref[m], 0, max(7, last_col - 5 + 3 * k), y_off=4)
    hdr = ["Region", "Segment", "Sub-segment", "Sub-segment 1", "Sub-segment 2", "Unit"]
    for j, h in enumerate(hdr):
        ms.write_string(5, j, h, f_h_txt)
    for j, y in enumerate(ds.years):
        ms.write_string(5, 6 + j, str(y), f_h_yr)
    r = 6
    idx = ds.index()
    for g in ds.geo_order:
        if g not in idx:
            continue
        for m in METRICS:
            by_s = idx[g].get(m)
            if not by_s:
                continue
            for s in ds.segment_order(g, list(by_s)):
                for a, b, c, v in by_s[s]:
                    for j, t in enumerate((g, s, a, b, c, unit_lbl[m])):
                        ms.write_string(r, j, t, f_txt)
                    for j, x in enumerate(_round_vec(v, dec[m])):
                        if isinstance(x, float):
                            ms.write_number(r, 6 + j, x, f_num)
                        elif x is None:
                            ms.write_blank(r, 6 + j, None, f_num)
                        else:
                            ms.write_string(r, 6 + j, str(x), f_num)
                    r += 1
    ms.autofilter(5, 0, max(r - 1, 6), last_col)
    master_rows = r - 6

    # ---- Value / Volume ----
    pivot_counts = {}
    for m in present:
        ws = wb.add_worksheet(m)
        ws.hide_gridlines(2)
        ws.set_zoom(ZOOM)
        ws.set_column(0, 0, PIVOT_LABEL_WIDTH)
        ws.set_column(1, n, PIVOT_YEAR_WIDTH)
        ws.merge_range(0, 0, 3, n, _title(ds, [m], top), f_title)
        button(ws, "Home", ref["Home"], 0, 0, x_off=151, y_off=4)
        nav = ["Master Sheet"] + [o for o in present if o != m]
        for k, t in enumerate(nav):
            button(ws, t, ref[t], 0, max(2, n - 4 + 2 * k), y_off=4)
        ws.write_string(18, 0, "Unit", f_plain)
        ws.write_string(18, 1, unit_lbl[m], f_plain)
        ws.write_string(20, 0, "Row Labels", f_plain)
        for j, y in enumerate(ds.years):
            ws.write_string(20, 1 + j, f" {y}", f_plain)
        rows, grand = _pivot_rows(ds, m, dec[m])
        rr = 21
        for lvl, label, vals in rows + [(0, "Grand Total", grand)]:
            ws.write_string(rr, 0, label, f_lbl[min(lvl, 5)])
            for j, x in enumerate(vals):
                if isinstance(x, float):
                    ws.write_number(rr, 1 + j, round(x, dec[m]), f_pnum[m])
                elif x is not None:
                    ws.write_string(rr, 1 + j, str(x), f_pnum[m])
            rr += 1
        pivot_counts[m] = len(rows) + 1
    wb.close()
    return {"master_rows": master_rows, "pivot_rows": pivot_counts, "top_geo": top,
            "geos": len(ds.geo_order), "years": f"{ds.years[0]}-{ds.years[-1]}"}


# --------------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------------
def default_out(ds, src):
    name = f"{ds.top_geo()} {ds.market or 'Market'} Market Formatted Dataset"
    name = re.sub(r"[^\w&\-]+", "_", name).strip("_")
    return os.path.join(os.path.dirname(os.path.abspath(src)), name + ".xlsx")


def main(argv=None):
    ap = argparse.ArgumentParser(description="Build the FMI Formatted Dataset (Home | Master Sheet | Value | Volume)")
    ap.add_argument("input", help="FMI Excel Report or Raw Data workbook (.xlsx)")
    ap.add_argument("-o", "--output", help="output .xlsx (default: <Geo>_<Market>_Market_Formatted_Dataset.xlsx)")
    ap.add_argument("--market", help="market name for titles, e.g. 'Packaged Muffins' (default: from the source title)")
    ap.add_argument("--unit-style", choices=["metric", "unit"], default="metric",
                    help="Unit column: 'metric' = Value/Volume (default), 'unit' = US$ Mn/MT")
    ap.add_argument("--value-unit", help="override value unit label in titles, e.g. 'US$ Mn'")
    ap.add_argument("--volume-unit", help="override volume unit label in titles, e.g. 'MT'")
    ap.add_argument("--value-decimals", type=int, default=3)
    ap.add_argument("--volume-decimals", type=int, default=3)
    ap.add_argument("--no-cross", action="store_true", help="skip the Cross-sectional sheet")
    ap.add_argument("--no-buttons", action="store_true", help="skip navigation buttons")
    ap.add_argument("--strict", action="store_true", help="exit 1 if QC reports any ERROR")
    a = ap.parse_args(argv)

    ds = read_excel_report(a.input, market=a.market, include_cross=not a.no_cross) \
        if _is_report(a.input) else read_raw_data(a.input, market=a.market)
    if a.value_unit:
        ds.units[VALUE] = a.value_unit
    if a.volume_unit:
        ds.units[VOLUME] = a.volume_unit
    issues = run_qc(ds, a.value_decimals, a.volume_decimals)
    out = a.output or default_out(ds, a.input)
    info = write_formatted_dataset(ds, out, a.unit_style, a.value_decimals, a.volume_decimals,
                                   buttons=not a.no_buttons)
    errs = [i for i in issues if i[0] == "ERROR"]
    print(f"Wrote {out}")
    print(f"  top geography: {info['top_geo']} | geographies: {info['geos']} | years: {info['years']}")
    print(f"  Master Sheet rows: {info['master_rows']:,} | pivot rows: "
          + ", ".join(f"{k} {v:,}" for k, v in info['pivot_rows'].items()))
    print(f"  units: Value = {ds.units.get(VALUE)}, Volume = {ds.units.get(VOLUME)}")
    print(f"QC: {len(errs)} error(s), {sum(1 for i in issues if i[0] == 'WARN')} warning(s)")
    for lvl, msg in issues[:60]:
        print(f"  [{lvl}] {msg}")
    if len(issues) > 60:
        print(f"  ... {len(issues) - 60} more")
    return 1 if (a.strict and errs) else 0


def _is_report(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        for name in wb.sheetnames:
            rows = list(wb[name].iter_rows(max_row=40, values_only=True))
            if _header(rows, ("geography", "segment", "metric"))[0] is not None:
                return True
        return False
    finally:
        wb.close()


if __name__ == "__main__":
    sys.exit(main())
